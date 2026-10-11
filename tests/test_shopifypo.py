import asyncio
import csv
import json
import re
from pathlib import Path

import httpx
import pytest

from agentkit.inbox import read_inbox
from agentkit.llm import LLM
from agentkit.runner import Runner
from shopifypo.catalog import Catalog, find_customer, read_catalog, read_customers, read_prices
from shopifypo.check import check_line
from shopifypo.extract import POLine
from shopifypo.output import write_outputs
from shopifypo.pipeline import CACHE_VERSION, read_all, read_seen
from shopifypo.push import Shopify, push
from shopifypo.sample import make_sample, score
from shopifypo.shop import ShopConfig, load_shop

from .fakes import FakeClient, no_sleep, user_prompt

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "shopifypo"


@pytest.fixture(scope="module")
def case(tmp_path_factory):
    out = tmp_path_factory.mktemp("shopifypo")
    truth = make_sample(out)
    variants, problems = read_catalog(out / "catalog.csv")
    catalog = Catalog(variants, read_prices(out / "prices.csv"))
    customers = read_customers(out / "customers.csv")
    return out, truth, catalog, customers, load_shop(out / "shop.toml")


def oracle(truth: dict):
    """A perfect stand-in model: answers with the answer key's reading of each email."""
    by_file = {e["file"]: e["reading"] for e in truth["emails"]}

    def respond(kwargs):
        return json.dumps(by_file[re.search(r"^Email file: (.+)$", user_prompt(kwargs), re.M).group(1)])
    return respond


def _run(case, tmp_path, responder=None, seen=None):
    out, truth, catalog, customers, cfg = case
    emails, _ = read_inbox(out / "inbox")
    runner = Runner(LLM(FakeClient(responder or oracle(truth)), "test-model", rpm=0, sleep=no_sleep),
                    tmp_path / "cache", CACHE_VERSION)
    seen = {} if seen is None else seen
    orders = asyncio.run(read_all(emails, catalog, customers, cfg, runner, seen, progress=lambda _: None))
    return orders, runner, seen


# --- catalog, prices and customers --------------------------------------------------------------------------------

def test_catalog_reads_shopify_export(case):
    out, _, catalog, _, _ = case
    variants, problems = read_catalog(out / "catalog.csv")
    assert len(variants) == 12 and problems == []  # the extra image row is skipped
    cs12 = next(v for v in variants if v.sku == "JP-CS-12")
    assert (cs12.title, cs12.options, cs12.price, cs12.inventory) == ("Cedar & Sage Soy Candle", [("Size", "12 oz")],
                                                                     34.0, 90)
    match = next(v for v in variants if v.sku == "JP-MATCH")
    assert match.options == [] and match.label == "Safety Matches"
    assert next(v for v in variants if v.sku == "JP-GIFT-24").status == "archived"
    assert catalog.generic == {"soy", "candle"}


def test_lines_are_matched_in_code(case):
    _, _, catalog, customers, _ = case
    willow = catalog.customer_rows(customers[1], "buyer@willowstreetgifts.example")
    m = catalog.match

    def sku(*a):
        r = m(*a)
        return (r.variant.sku if r.variant else None, r.how)
    assert sku("jp cs 08", "", "", "", []) == ("JP-CS-08", "SKU")
    assert sku("", "WSG-1012", "", "Cedar Sage 12oz candle", willow) == ("JP-CS-12", "buyer's SKU")
    assert sku("", "", "0850012340097", "", []) == ("JP-WT-BR", "barcode")  # EAN-13 of the UPC
    assert sku("", "", "", "Fig Leaf JP-FL-12 large", []) == ("JP-FL-12", "SKU in description")
    assert sku("", "", "", "Cedar Sage 12oz candle", []) == ("JP-CS-12", "description")
    assert sku("", "", "", "wick trimmer, black finish", []) == ("JP-WT-BK", "description")
    amb = m("", "", "", "Sea Salt Linen candle", [])
    assert amb.variant is None and [v.sku for v in amb.candidates] == ["JP-SL-08", "JP-SL-12"]
    assert m("", "", "", "Lavender Soy Candle 8 oz", []).variant is None


def test_expected_prices_and_customers(case):
    _, _, catalog, customers, cfg = case
    by = {v.sku: v for v in catalog.variants}
    harbor = catalog.customer_rows(customers[0], "orders@harborvine.example")
    assert catalog.expected(by["JP-CS-08"], harbor, cfg) == (11.5, "this buyer's price list", 6, 0)
    assert catalog.expected(by["JP-CS-08"], [], cfg) == (12.0, "price list", 6, 0)
    assert catalog.expected(by["JP-GIFT-24"], [], cfg) == (29.0, "retail $58.00 less 50%", 0, 0)
    assert find_customer("ORDERS@harborvine.example", "", customers)[1] == "email"
    assert find_customer("raj@willowstreetgifts.example", "", customers)[1] == "email domain"
    assert find_customer("x@gmail.com", "The Paper Lantern, LLC", customers)[1] == "company name"
    assert find_customer("x@gmail.com", "Other Co", customers) == (None, "")


def test_line_checks(case):
    _, _, catalog, _, cfg = case

    def check(text, **kw):
        ln = check_line(1, POLine(**kw), text, catalog, [], cfg)
        return ln, {f.code: f.hold for f in ln.flags}

    ln, codes = check("JP-AS-08 Amber 24 @ 10.00 = 240.00", vendor_sku="JP-AS-08", quantity=24, unit_price=10,
                      line_total=240)
    assert codes == {"price_mismatch": True} and ln.unit_price == 12.0
    ln, codes = check("JP-CS-08 24 x 12.00 = 268.00", vendor_sku="JP-CS-08", quantity=24, unit_price=12, line_total=268)
    assert codes == {"line_math": True}
    ln, codes = check("2 cases JP-FL-08 at 72.00 a case", vendor_sku="JP-FL-08", quantity=2, unit="case",
                      unit_price=72)
    assert (ln.units, ln.po_price, ln.unit_price) == (12, 12.0, 12.0) and codes == {"case_converted": False}
    ln, codes = check("JP-MATCH x 6", vendor_sku="JP-MATCH", quantity=6)
    assert codes == {"case_pack": True, "no_po_price": False}
    ln, codes = check("JP-AS-12 x 12", vendor_sku="JP-AS-12", quantity=12)
    assert codes == {"short_stock": False, "no_po_price": False}
    ln, codes = check("JP-SL-08 x 12", vendor_sku="JP-SL-09", quantity=18, unit_price=12)
    assert codes["sku_not_in_text"] is False and codes["qty_not_in_text"] is True and codes["unmatched"] is True


def test_shop_settings(tmp_path):
    p = tmp_path / "s.toml"
    p.write_text("[shop]\nname = 'X'\nwholesale_discount = 50\n")
    with pytest.raises(ValueError, match="share of the retail price"):
        load_shop(p)
    p.write_text("[shop]\nmargin = 1\n")
    with pytest.raises(ValueError, match="unknown \\[shop\\] settings: margin"):
        load_shop(p)


# --- the whole example --------------------------------------------------------------------------------------------

def test_committed_example_matches_the_generator(tmp_path):
    truth = make_sample(tmp_path)
    assert truth == json.loads((EXAMPLE / "truth.json").read_text())
    for f in [*(tmp_path / "inbox").iterdir(), *(p for p in tmp_path.iterdir() if p.is_file())]:
        rel = f.relative_to(tmp_path)
        assert f.read_bytes() == (EXAMPLE / rel).read_bytes(), rel


def test_end_to_end_on_the_synthetic_inbox_scores_full_marks(case, tmp_path):
    _, truth, catalog, customers, cfg = case
    orders, runner, seen = _run(case, tmp_path)
    assert runner.llm.usage.requests == 7
    write_outputs(tmp_path, orders, cfg.tags, seen)
    s = score(json.loads((tmp_path / "orders.json").read_text()), truth)
    assert (s["kind_right"], s["ready_right"], s["customer_right"]) == (7, 7, 7)
    assert (s["lines_found"], s["lines"], s["missed_lines"]) == (18, 18, [])
    assert (s["price_right"], s["wrong_prices"]) == (18, [])
    assert (s["flags_found"], s["flags_expected"], s["extra_holds"]) == (11, 11, [])

    drafts = sorted(p.name for p in (tmp_path / "draft_orders").iterdir())
    assert drafts == ["HV_2291.json"]
    d = json.loads((tmp_path / "draft_orders" / "HV_2291.json").read_text())
    assert d["email"] == "orders@harborvine.example" and d["expected_subtotal"] == 828.0
    assert d["lines"][0] == {"sku": "JP-CS-08", "title": "Cedar & Sage Soy Candle - 8 oz", "quantity": 24,
                             "unit_price": 11.5}
    assert d["ship_to"]["zip"] == "04101" and "terms Net 30" in d["note"]
    with open(tmp_path / "review.csv") as f:
        review = list(csv.DictReader(f))
    assert review[0]["must_fix"] == "yes" and not any(r["email"] == "01_harbor_vine_po.eml" for r in review)
    dup = next(r for r in review if r["email"] == "05_harbor_vine_resend.eml")
    assert dup["check"].startswith("PO HV-2291 was already received (01_harbor_vine_po.eml, 2026-10-05)")
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "orders.xlsx")
    assert wb.sheetnames == ["Orders", "Lines", "Review"] and wb["Orders"].cell(2, 5).value == "ready"

    # the second run comes from the cache, and the PO log catches the same POs sent again
    again = Runner(LLM(FakeClient(lambda kw: 1 / 0), "test-model", rpm=0, sleep=no_sleep), tmp_path / "cache",
                   CACHE_VERSION)
    emails, _ = read_inbox(case[0] / "inbox")
    later = asyncio.run(read_all(emails[:1], catalog, customers, cfg, again, read_seen(tmp_path / "po_log.csv"),
                                 progress=lambda _: None))
    assert again.from_cache == 1 and [f.code for f in later[0].flags] == ["duplicate"]


def test_the_model_over_reaching_is_caught(case, tmp_path):
    base = oracle(case[1])

    def sloppy(kwargs):
        reading = json.loads(base(kwargs))
        prompt = user_prompt(kwargs)
        if "Email file: 01_harbor_vine_po.eml" in prompt:
            reading["lines"][0]["quantity"] = 36  # misread
            reading["lines"][1]["vendor_sku"] = "JP-FL-12"  # not what the PO says
        if "Email file: 03_paper_lantern.txt" in prompt:
            reading["po_number"] = "PL-7781"  # invented format
        return json.dumps(reading)

    orders, _, _ = _run(case, tmp_path, sloppy)
    hv = orders[0]
    assert not hv.ready
    assert any(f.code == "qty_not_in_text" and f.hold for f in hv.lines[0].flags)
    assert hv.lines[1].sku == "JP-FL-08" and hv.lines[1].how == "description"  # the invented SKU was ignored
    assert [f.code for f in orders[2].flags if f.hold] == ["po_number_not_in_text", "po_number_missing"]


# --- pushing to Shopify -------------------------------------------------------------------------------------------

def _fake_shopify(calls: list):
    variants = {"JP-CS-08": ("gid://shopify/ProductVariant/1", "24.00"), "JP-WT-BR": ("gid://shopify/ProductVariant/2",
                                                                                       "18.00")}

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["X-Shopify-Access-Token"] == "tok"
        body = json.loads(request.content)
        calls.append(body)
        if "productVariants" in body["query"]:
            sku = re.search(r'sku:"(.+)"', body["variables"]["q"]).group(1)
            edges = [{"node": {"id": variants[sku][0], "sku": sku, "price": variants[sku][1],
                               "product": {"title": "x", "status": "ACTIVE"}}}] if sku in variants else []
            return httpx.Response(200, json={"data": {"productVariants": {"edges": edges}}})
        lines = body["variables"]["input"]["lineItems"]
        sub = sum(float(variants["JP-CS-08" if "1" in x["variantId"] else "JP-WT-BR"][1]) * x["quantity"]
                  * (1 - x.get("appliedDiscount", {}).get("value", 0) / 100) for x in lines)
        return httpx.Response(200, json={"data": {"draftOrderCreate": {
            "draftOrder": {"id": "gid://shopify/DraftOrder/9", "name": "#D9",
                           "subtotalPriceSet": {"shopMoney": {"amount": f"{sub:.2f}"}}}, "userErrors": []}}})
    return httpx.Client(transport=httpx.MockTransport(handler))


DRAFT = {"file": "po.eml", "po_number": "HV-1", "email": "orders@harborvine.example", "company": "Harbor & Vine",
         "ship_to": {"name": "Elena Brooks", "company": "Harbor & Vine", "address1": "88 Commercial St",
                     "city": "Portland", "state": "ME", "zip": "04101", "country": "US"},
         "tags": ["wholesale"], "note": "PO HV-1",
         "lines": [{"sku": "JP-CS-08", "title": "x", "quantity": 24, "unit_price": 11.5},
                   {"sku": "JP-WT-BR", "title": "y", "quantity": 8, "unit_price": 9.0}], "expected_subtotal": 348.0}


def test_push_is_a_dry_run_until_asked_and_checks_shopifys_total():
    calls = []
    shop = Shopify("store.myshopify.com", "tok", "2025-07", client=_fake_shopify(calls))
    plans = push([DRAFT], shop, yes=False)
    assert plans[0].created == "" and plans[0].problems == [] and len(calls) == 2  # two lookups, nothing created
    item = plans[0].draft_input["lineItems"][0]
    assert item["appliedDiscount"]["value"] == pytest.approx(52.0833) and item["quantity"] == 24
    assert plans[0].draft_input["shippingAddress"]["provinceCode"] == "ME"
    plans = push([DRAFT], shop, yes=True)
    assert plans[0].created == "#D9" and plans[0].shopify_subtotal == pytest.approx(348.0) and plans[0].problems == []
    bad = {**DRAFT, "lines": [{"sku": "JP-CS-08", "title": "x", "quantity": 6, "unit_price": 30.0},
                              {"sku": "JP-NOPE", "title": "z", "quantity": 1, "unit_price": 1.0}]}
    plans = push([bad], shop, yes=True)
    assert plans[0].created == "" and plans[0].problems == [
        "JP-CS-08: wholesale $30.00 is above the store price $24.00", "JP-NOPE not found in Shopify (or not unique)"]


# --- command line -------------------------------------------------------------------------------------------------

def test_cli_run_score_and_push(tmp_path, monkeypatch, capsys):
    from shopifypo import __main__ as cli

    ex = tmp_path / "ex"
    cli.main(["sample", "--out", str(ex)])
    truth = json.loads((ex / "truth.json").read_text())
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(cli.LLM, "from_settings",
                        classmethod(lambda cls, s: LLM(FakeClient(oracle(truth)), s.model, rpm=0, sleep=no_sleep)))
    out = tmp_path / "out"
    cli.main(["run", "--emails", str(ex / "inbox"), "--catalog", str(ex / "catalog.csv"), "--prices",
              str(ex / "prices.csv"), "--customers", str(ex / "customers.csv"), "--shop", str(ex / "shop.toml"),
              "--out", str(out), "--model", "x:free"])
    printed = capsys.readouterr().out
    assert "WARNING: free models may log prompts" in printed
    assert "Done: 6 purchase orders, 1 ready for Shopify, 6 to fix, 0 not a PO; 7 model requests" in printed
    cli.main(["score", "--run", str(out), "--truth", str(ex / "truth.json")])
    report = capsys.readouterr().out
    assert "Lines matched to the right SKU and quantity: 18 of 18 (100%)" in report
    assert "Expected checks raised: 11 of 11 (100%); other must-fix checks: 0" in report
    monkeypatch.delenv("SHOPIFY_ACCESS_TOKEN", raising=False)
    with pytest.raises(SystemExit, match="SHOPIFY_ACCESS_TOKEN"):
        cli.main(["push", "--run", str(out), "--shop", str(ex / "shop.toml")])
