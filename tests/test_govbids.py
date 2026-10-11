import asyncio
import csv
import json
import re
from datetime import date
from pathlib import Path

import httpx
import pytest

from agentkit.docs import read_document
from agentkit.llm import LLM
from agentkit.runner import Runner
from govbids.fetch import fetch_notices
from govbids.match import NoticeScore, check_score
from govbids.notices import Notice, apply_rules, read_notices, required_clearance, set_aside_code
from govbids.output import write_matches, write_shred
from govbids.pipeline import CACHE_VERSION, match_all, shred
from govbids.profile import load_profile
from govbids.sample import make_sample, score
from govbids.shred import Facts, check_facts, find_requirements

from .fakes import FakeClient, no_sleep, user_prompt

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "govbids"
TODAY = date(2026, 10, 12)


@pytest.fixture(scope="module")
def case(tmp_path_factory):
    out = tmp_path_factory.mktemp("govbids")
    truth = make_sample(out)
    return out, truth, load_profile(out / "company.toml")


def oracle(truth: dict, scores=None, facts=None):
    scores = scores or truth["scores"]

    def respond(kwargs):
        prompt = user_prompt(kwargs)
        if prompt.startswith("OUR COMPANY"):
            ids = re.findall(r"^### (\S+)$", prompt, re.M)
            return json.dumps({"notices": [{"id": i, **scores[i]} for i in ids]})
        if prompt.startswith("PROPOSAL VOLUMES"):
            out = []
            for rid, text in re.findall(r"^\[(R\d+)\] \([^)]*\) (.*)$", prompt, re.M):
                lab = next(x for x in truth["requirements"] if x["phrase"].lower() in text.lower())
                out.append({"id": rid, "type": lab["type"], "volume": lab["volume"], "summary": lab["phrase"],
                            "offeror": lab["offeror"]})
            return json.dumps({"requirements": out})
        return json.dumps(facts or truth["facts"])
    return respond


def _runner(tmp_path, responder):
    return Runner(LLM(FakeClient(responder), "test-model", rpm=0, sleep=no_sleep), tmp_path / "cache", CACHE_VERSION)


# --- profile and rules --------------------------------------------------------------------------------------------

def test_profile_settings(case, tmp_path):
    _, _, p = case
    assert p.certifications == ["SB", "SDVOSB"] and p.places == ["VA", "MD", "DC"] and p.clearance == "secret"
    bad = tmp_path / "c.toml"
    bad.write_text('[company]\nname = "X"\nnaics = ["541512"]\ncertifications = ["8(a)", "HUBZONE"]\n')
    with pytest.raises(ValueError, match="unknown certifications \\['HUBZONE'\\]"):
        load_profile(bad)


def test_notice_rules_in_code(case):
    out, truth, p = case
    assert set_aside_code("", "Service-Disabled Veteran-Owned Small Business (SDVOSB) Set-Aside") == "SDVOSBC"
    assert set_aside_code("", "8(a) Sole Source") == "8A" and set_aside_code("SBA", "") == "SBA"
    assert required_clearance("staff must hold an active TS/SCI clearance") == "ts/sci"
    assert required_clearance("requires a security clearance at the Secret level") == "secret"
    assert required_clearance("no clearance needed") == ""
    notices = read_notices(out / "notices.csv")
    for n in notices:
        apply_rules(n, p, TODAY)
    by = {n.id: n for n in notices}
    for t in truth["notices"]:
        n = by[t["id"]]
        if t["status"] == "filtered out":
            assert t["filtered_reason"] in "; ".join(n.out), n.id
        else:
            assert n.out == [], n.id
    assert by["EXA-HD-9920"].notes == ["requires a TS/SCI clearance; we have secret"]
    assert by["ELC-HD-2026-014"].deadline == date(2026, 11, 6)


def test_requirements_are_found_word_for_word(case):
    out, truth, _ = case
    reqs = find_requirements(read_document(out / "rfp" / "RFP_ELC-HD-2026-014.pdf"))
    assert len(reqs) == 20 and [r.section for r in reqs[:3]] == ["C.1", "C.2", "C.3"]
    assert reqs[8].text == ("Questions must be submitted by email to the Contracting Officer no later than October "
                            "23, 2026.")
    assert reqs[8].source == "RFP_ELC-HD-2026-014.pdf p.2"
    assert not any("are due no later" in r.text for r in reqs)  # a key fact, not a requirement


def test_score_checks_quotes_and_decisions(case):
    _, _, p = case
    n = Notice(id="X", title="Sources Sought: ServiceNow", type="Sources Sought",
               description="We seek firms that administer ServiceNow.")
    m = check_score(NoticeScore.model_validate({"id": "X", "fit": 90, "decision": "bid", "matches": [
        {"capability": "ServiceNow", "quote": "We need a ServiceNow upgrade by May"}]}), n)
    assert m.fit == 40 and m.decision == "respond"
    assert m.checks == ["1 quotes dropped: not in the notice",
                        "fit lowered from 90 to 40: no capability quoted from the notice",
                        "a sources sought can't be bid on yet; marked 'respond'",
                        "respond with a fit of 40: check it's worth the effort"]


def test_key_facts_need_matching_quotes():
    text = "Proposals are due no later than November 6, 2026 at 2:00 PM. Volume I shall not exceed 20 pages."
    rows, checks = check_facts(Facts.model_validate({
        "due": {"date": "2026-11-07", "quote": "Proposals are due no later than November 6, 2026 at 2:00 PM."},
        "page_limits": [{"volume": "Volume I", "pages": 25, "quote": "Volume I shall not exceed 20 pages."}],
        "submission": {"value": "email", "quote": "Send it by email to the CO"}}), text)
    assert checks == ["Proposal due: the date 2026-11-07 is not the one in its quote",
                      "Submission dropped: its quote is not in the RFP",
                      "Volume I page limit 25 is not the number in its quote"]
    assert [r["fact"] for r in rows] == ["Proposal due", "Questions due"] and rows[0]["value"] == ""


# --- the whole example --------------------------------------------------------------------------------------------

def test_committed_example_matches_the_generator(tmp_path):
    truth = make_sample(tmp_path)
    assert truth == json.loads((EXAMPLE / "truth.json").read_text())
    for f in tmp_path.rglob("*"):
        if f.is_file():
            rel = f.relative_to(tmp_path)
            assert f.read_bytes() == (EXAMPLE / rel).read_bytes(), rel


def test_end_to_end_match_and_shred_score_full_marks(case, tmp_path):
    out, truth, p = case
    runner = _runner(tmp_path, oracle(truth))
    matches = asyncio.run(match_all(read_notices(out / "notices.csv"), p, runner, TODAY, progress=lambda _: None))
    assert runner.llm.usage.requests == 1  # 5 notices left after the rules, one batch
    write_matches(tmp_path, matches, p.name, TODAY)
    reqs, facts, checks = asyncio.run(shred(read_document(out / "rfp" / "RFP_ELC-HD-2026-014.pdf"), runner,
                                            progress=lambda _: None))
    assert runner.llm.usage.requests == 3 and checks == []
    write_shred(tmp_path, reqs, facts, checks)
    s = score(json.loads((tmp_path / "matches.json").read_text()), json.loads((tmp_path / "shred.json").read_text()),
              truth)
    assert (s["status_right"], s["notices"], s["reason_right"], s["reason_expected"]) == (12, 12, 7, 7)
    assert (s["found"], s["requirements"], s["extra"], s["labels_right"]) == (20, 20, 0, 20)
    assert (s["facts_found"], s["facts_expected"]) == (9, 9)

    digest = (tmp_path / "digest.md").read_text()
    assert digest.startswith("# Cobalt Ridge Systems LLC: SAM.gov notices, October 12, 2026")
    assert "## BID: IT Help Desk and End User Support Services (fit 92)" in digest
    assert "## RESPOND: Sources Sought: ServiceNow Platform Administration (fit 85)" in digest
    assert "Also read: 2 scored as a pass; 7 filtered out by the rules" in digest
    with open(tmp_path / "compliance_matrix.csv") as f:
        rows = list(csv.DictReader(f))
    assert rows[7]["offeror_must"] == "no" and rows[0]["volume"] == "Volume I - Technical"
    with open(tmp_path / "key_facts.csv") as f:
        kf = {r["fact"]: r["value"] for r in csv.DictReader(f)}
    assert kf["Proposal due"] == "2026-11-06 2:00 PM Eastern Time" and kf["Page limit: Volume I - Technical"] == "20 pages"
    from openpyxl import load_workbook
    assert load_workbook(tmp_path / "compliance_matrix.xlsx").sheetnames == ["Compliance matrix", "Key facts"]


# --- SAM.gov ------------------------------------------------------------------------------------------------------

def test_fetch_reads_the_sam_api(tmp_path):
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.params["api_key"] == "k"
        if request.url.path.endswith("/search"):
            assert request.url.params["postedFrom"] == "10/05/2026" and request.url.params["ncode"] == "541512"
            return httpx.Response(200, json={"totalRecords": 1, "opportunitiesData": [{
                "noticeId": "abc", "title": "Help desk", "solicitationNumber": "S-1", "fullParentPathName": "DEPT.X",
                "postedDate": "2026-10-06", "type": "Solicitation", "typeOfSetAside": "SBA",
                "typeOfSetAsideDescription": "Total Small Business Set-Aside (FAR 19.5)",
                "responseDeadLine": "2026-11-01T12:00:00-04:00", "naicsCode": "541512",
                "placeOfPerformance": {"city": {"name": "Reston"}, "state": {"code": "VA"}},
                "uiLink": "https://sam.example/opp/abc", "description": "https://api.sam.example/desc?noticeid=abc"}]})
        return httpx.Response(200, json={"description": "<p>Tier 1 help desk &amp; ServiceNow.</p>"})

    rows = fetch_notices("k", date(2026, 10, 5), date(2026, 10, 12), ["541512"],
                         client=httpx.Client(transport=httpx.MockTransport(handler)))
    assert rows[0]["PopState"] == "VA" and rows[0]["Description"] == "Tier 1 help desk & ServiceNow."
    from govbids.fetch import save
    save(rows, tmp_path / "n.csv")
    n = read_notices(tmp_path / "n.csv")[0]
    assert (n.id, n.set_aside_code, n.deadline) == ("abc", "SBA", date(2026, 11, 1))


# --- command line -------------------------------------------------------------------------------------------------

def test_cli_notices_match_shred_and_score(tmp_path, monkeypatch, capsys):
    from govbids import __main__ as cli

    ex = tmp_path / "ex"
    cli.main(["sample", "--out", str(ex)])
    truth = json.loads((ex / "truth.json").read_text())
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(cli.LLM, "from_settings",
                        classmethod(lambda cls, s: LLM(FakeClient(oracle(truth)), s.model, rpm=0, sleep=no_sleep)))
    base = ["--company", str(ex / "company.toml"), "--notices", str(ex / "notices.csv"), "--today", "2026-10-12"]
    cli.main(["notices", *base])
    assert "12 notices; 5 left to score = 1 model request." in capsys.readouterr().out
    out = tmp_path / "out"
    cli.main(["match", *base, "--out", str(out)])
    assert "Done: 2 to bid, 1 to respond to, 2 to pass, 7 filtered out; 1 model request" in capsys.readouterr().out
    cli.main(["shred", "--rfp", str(ex / "rfp"), "--out", str(out)])
    assert "Done: 20 requirements (16 for the offeror), 9 key facts; 2 model requests" in capsys.readouterr().out
    cli.main(["score", "--run", str(out), "--truth", str(ex / "truth.json")])
    report = capsys.readouterr().out
    assert "Notices sorted right (bid / respond / pass / filtered out): 12 of 12 (100%)" in report
    assert "RFP requirements found: 20 of 20 (100%); other sentences: 0" in report
