import asyncio
import json
from datetime import date
from email.message import EmailMessage

import httpx
import openai
import pytest
from pydantic import BaseModel

from agentkit.checks import (amounts, emails_in, find_dates, find_phrases, parse_date, parse_money, quote_found,
                             relative_dates, unsupported_numbers, valid_email)
from agentkit.inbox import read_inbox
from agentkit.llm import LLM
from agentkit.runner import NOT_YET, Runner
from agentkit.settings import Settings, is_free_model, plural
from agentkit.tables import Sheet, pick, read_table, sheet_text, write_xlsx

from .fakes import FakeClient, no_sleep


class Out(BaseModel):
    ok: bool = False


def test_settings_prefer_the_agents_own_variables(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "k1")
    monkeypatch.setenv("LEADAGENT_MODEL", "base/model:free")
    monkeypatch.setenv("LEADAGENT_FALLBACK_MODELS", "a:free, b:free")
    monkeypatch.setenv("OUTREACH_MODEL", "paid/model")
    s = Settings.from_env("OUTREACH")
    assert (s.api_key, s.model, s.fallback_models, s.app_title) == ("k1", "paid/model", ["a:free", "b:free"], "outreach")
    assert Settings.from_env("GRANTS").model == "base/model:free"
    assert is_free_model("openrouter/free") and is_free_model("x/y:free") and not is_free_model("x/y")
    assert plural(1, "reply", "replies") == "1 reply" and plural(2, "reply", "replies") == "2 replies"


def _rate_limit(msg: str) -> openai.RateLimitError:
    req = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions")
    return openai.RateLimitError(msg, response=httpx.Response(429, request=req), body=None)


def test_runner_caches_and_stops_at_the_daily_limit(tmp_path):
    calls = []

    def respond(kw):
        calls.append(1)
        if len(calls) > 1:
            raise _rate_limit("Rate limit exceeded: free-models-per-day")
        return '{"ok": true}'

    r = Runner(LLM(FakeClient(respond), "m", rpm=0, sleep=no_sleep), tmp_path, "1")
    got, err = asyncio.run(r.ask("sys", "first", Out))
    assert got.ok and err == ""
    got, err = asyncio.run(r.ask("sys", "second", Out))
    assert got is None and err == NOT_YET and "daily request limit" in r.stopped_reason
    got, err = asyncio.run(r.ask("sys", "third", Out))
    assert (got, err, len(calls)) == (None, NOT_YET, 2)  # no request once stopped

    again = Runner(LLM(FakeClient(lambda kw: 1 / 0), "m", rpm=0, sleep=no_sleep), tmp_path, "1")
    got, _ = asyncio.run(again.ask("sys", "first", Out))
    assert got.ok and again.from_cache == 1
    newer = Runner(LLM(FakeClient(lambda kw: '{"ok": false}'), "m", rpm=0, sleep=no_sleep), tmp_path, "2")
    got, _ = asyncio.run(newer.ask("sys", "first", Out))
    assert not got.ok and newer.from_cache == 0  # a new CACHE_VERSION asks again


def test_runner_stops_after_the_model_stays_busy(tmp_path):
    def busy(kw):
        raise _rate_limit("Provider returned error: busy")

    llm = LLM(FakeClient(busy), "m", rpm=0, sleep=no_sleep, max_retries=0)
    r = Runner(llm, tmp_path, "1", busy_limit=2, hint="set X_FALLBACK_MODELS")
    _, e1 = asyncio.run(r.ask("s", "a", Out))
    _, e2 = asyncio.run(r.ask("s", "b", Out))
    _, e3 = asyncio.run(r.ask("s", "c", Out))
    assert e1.startswith("model request failed (RateLimitError") and r.errors == 2
    assert "busy for 2 requests in a row" in r.stopped_reason and "X_FALLBACK_MODELS" in r.stopped_reason
    assert e3 == NOT_YET


def test_numbers_must_be_in_the_source_as_whole_values():
    assert unsupported_numbers("kept for 3 months", "kept for 35 days") == ["3"]
    assert unsupported_numbers("a $5M policy", "a $5,000,000 policy") == []
    assert unsupported_numbers("a $5M policy", "5 staff use AES-256") == ["5M"]
    assert unsupported_numbers("TLS 1.2, AES-256, 72 hours", "TLS 1.2 or higher; AES-256; within 72 hours") == []
    assert unsupported_numbers("we reached 1,250 students", "1250 students") == []
    assert amounts("twelve schools, 25% growth, $4.5M") >= {12.0, 25.0, 4.5, 4_500_000.0}
    assert parse_money("USD 1,250.50") == 1250.5 and parse_money("n/a") is None


def test_dates_in_the_formats_letters_and_contracts_use():
    text = "Accepted 10/02/2026. Closing on November 20, 2026 or the 3rd day of December, 2026; ISO 2026-10-15."
    assert find_dates(text) == {date(2026, 10, 2), date(2026, 11, 20), date(2026, 12, 3), date(2026, 10, 15)}
    assert find_dates("due Oct 9 and 10/30") == set()  # no year and no reference date
    assert find_dates("due Oct 9 and 10/30", ref=date(2026, 10, 1)) == {date(2026, 10, 9), date(2026, 10, 30)}
    assert find_dates("back Jan 4", ref=date(2026, 12, 20)) == {date(2027, 1, 4)}
    assert parse_date("Sept. 15th, 2026") == date(2026, 9, 15) and parse_date("soon") is None
    wed = date(2026, 10, 7)
    assert relative_dates("back Monday", wed) == {date(2026, 10, 12)}
    assert relative_dates("back next Monday", wed) == {date(2026, 10, 12), date(2026, 10, 19)}
    assert relative_dates("call me tomorrow", wed) == {date(2026, 10, 8)}


def test_quotes_phrases_and_emails():
    src = "We notify affected customers within 72 hours of confirming a breach."
    assert quote_found("we notify affected customers within 72 hours", src)
    assert not quote_found("We notify customers within 48 hours", src)
    assert not quote_found("Our SOC 2 report covers availability", src)
    assert find_phrases("I hope this email finds you well! Quick question.", ["hope this email finds you well",
                                                                             "quick question", "synergy"]) == [
        "hope this email finds you well", "quick question"]
    assert emails_in("Write to Dana.Ruiz@acme.example. Or ops@acme.example") == ["dana.ruiz@acme.example",
                                                                                 "ops@acme.example"]
    assert valid_email("a@b.co") and not valid_email("a@b") and not valid_email("a..b@c.co")


def test_tables_find_the_header_and_read_csv_and_excel(tmp_path):
    (tmp_path / "t.csv").write_text("Export of products\n\nVariant SKU,Title,Variant Price\nJP-1,Candle,12.5\n,,\n")
    rows = read_table(tmp_path / "t.csv")
    assert rows == [{"variant sku": "JP-1", "title": "Candle", "variant price": "12.5"}]
    assert pick(rows[0], "SKU", "Variant_SKU") == "JP-1" and pick(rows[0], "barcode") == ""
    write_xlsx(tmp_path / "t.xlsx", [Sheet("Items", ["Variant SKU", "Qty"], [["JP-1", 3.0], ["JP-2", 4]], yellow={1})])
    assert read_table(tmp_path / "t.xlsx") == [{"variant sku": "JP-1", "qty": "3"}, {"variant sku": "JP-2", "qty": "4"}]
    assert sheet_text(tmp_path / "t.xlsx") == "Variant SKU | Qty\nJP-1 | 3\nJP-2 | 4"
    with pytest.raises(ValueError, match="save as .xlsx"):
        read_table(tmp_path / "old.xls")


def test_inbox_reads_excel_and_word_attachments(tmp_path):
    import docx

    write_xlsx(tmp_path / "po.xlsx", [Sheet("PO", ["SKU", "Qty"], [["JP-1", 24]])])
    d = docx.Document()
    d.add_paragraph("Ship to: 12 Harbor Rd")
    d.save(tmp_path / "note.docx")
    msg = EmailMessage()
    msg["From"], msg["Subject"], msg["Date"] = "Buyer <b@shop.example>", "PO 1001", "Tue, 6 Oct 2026 09:00:00 -0400"
    msg.set_content("PO attached.")
    msg.add_attachment((tmp_path / "po.xlsx").read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="po.xlsx")
    msg.add_attachment((tmp_path / "note.docx").read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="note.docx")
    mail = tmp_path / "mail"
    mail.mkdir()
    (mail / "po.eml").write_bytes(bytes(msg))
    emails, skipped = read_inbox(mail)
    assert skipped == [] and [a.name for a in emails[0].attachments] == ["po.xlsx", "note.docx"]
    assert "JP-1 | 24" in emails[0].all_text() and "Ship to: 12 Harbor Rd" in emails[0].all_text()


def test_documents_read_by_type(tmp_path):
    from agentkit.docs import document_files, read_document

    (tmp_path / "a.md").write_text("# Mission\nWe teach reading.\n")
    (tmp_path / "b.png").write_bytes(b"x")
    files, skipped = document_files(tmp_path)
    assert [f.name for f in files] == ["a.md"] and skipped == ["b.png (unsupported type)"]
    assert read_document(tmp_path / "a.md") == [("a.md § Mission", "We teach reading.")]
    with pytest.raises(ValueError, match="unsupported type"):
        read_document(tmp_path / "b.png")
    assert json.dumps(read_document(tmp_path / "a.md"))  # plain data
