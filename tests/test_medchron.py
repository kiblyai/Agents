import asyncio
import json
import re
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import httpx
import openai
import pytest

from agentkit.llm import LLM
from medchron.chronology import build, find_gaps, merge_bills, merge_entries, page_range
from medchron.extract import (Bill, CheckContext, Entry, ExtractedBill, ExtractedEntry, check_bill, check_entry,
                              find_dates, numbers, parse_date)
from medchron.pipeline import Options, extract_all, make_batches
from medchron.records import Page, ocr_pdf, read_case, write_combined
from medchron.sample import make_sample, score

from .fakes import FakeClient, no_sleep, user_prompt

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "medchron"
TODAY = date(2026, 9, 29)


# --- reading dates and numbers ------------------------------------------------------------------------------------

def test_dates_in_the_formats_records_use():
    text = ("DOS 03/14/2025, seen 3/24/25, April 15, 2025, 2025-04-22, 18-Jun-2025 and 9th July 2025. "
            "Phone 555-123-4567, SSN 123-45-6789, BP 138/86.")
    assert find_dates(text, TODAY) == {date(2025, 3, 14), date(2025, 3, 24), date(2025, 4, 15), date(2025, 4, 22),
                                       date(2025, 6, 18), date(2025, 7, 9)}
    assert find_dates("25/12/2024", TODAY) == {date(2024, 12, 25)}  # day first when it can only be day first
    assert parse_date("2025-03-14") == parse_date("03/14/2025") == date(2025, 3, 14)
    assert parse_date("unknown") is None


def test_service_date_labels_ignore_birth_statement_and_history_dates():
    from medchron.extract import service_dates

    text = ("Date of visit: April 15, 2025\nHistory: seen on 03/14/2025 in the ED. DOS 4/15/25\nDOB: 08/02/1986\n"
            "Statement date: 8/1/25\nDate: 3/27/25")
    assert service_dates(text, TODAY) == {date(2025, 4, 15), date(2025, 3, 27)}


def test_numbers_are_normalized_and_include_date_parts():
    assert {"1450", "7", "10", "3"} <= numbers("$1,450.00 pain 7/10, 03 visits")
    assert {"2025", "3", "17"} <= numbers("off work until March 17, 2025")


# --- checks on one extracted entry ----------------------------------------------------------------------------------

PAGES = {
    2: "Maple Falls General Hospital\nPatient: Dana Whitfield DOB: 08/02/1986\nDate of service: 03/14/2025\n"
       "Attending physician: Priya Raman, MD\nNeck pain 7/10.\nS13.4XXA Sprain of ligaments of cervical spine",
    3: "Maple Falls General Hospital\nPatient: Dana Whitfield\nIbuprofen 600 mg. Off work through 03/17/2025.\n"
       "03/14/2025 | 99284 | ED visit | 1 | $1,450.00",
    9: "Patient: Robert Hale\nDate: 4/8/25\nTherapist: Mark Chen, PT, DPT\nKnee pain 3/10",
}
CTX = CheckContext(patient="Dana Whitfield", doi=date(2025, 3, 14), today=TODAY)


def _entry(**kw) -> ExtractedEntry:
    base = dict(date="2025-03-14", provider="Priya Raman, MD", facility="Maple Falls General Hospital",
                visit_type="Emergency department visit", complaints="Neck pain 7/10",
                diagnoses=[{"code": "S13.4XXA", "description": "Sprain of ligaments of cervical spine"}],
                treatment="Ibuprofen 600 mg", work_status="Off work through 03/17/2025", pages=[2, 3],
                confidence="high")
    base.update(kw)
    return ExtractedEntry.model_validate(base)


def test_well_supported_entry_is_not_flagged():
    r = check_entry(_entry(), PAGES, {2, 3}, CTX)
    assert not r.needs_review and r.review_note == "" and r.pages == [2, 3] and r.diagnoses[0]["code"] == "S13.4XXA"


def test_wrong_date_invented_code_and_number_are_caught():
    r = check_entry(_entry(date="2025-03-15", treatment="Ibuprofen 800 mg",
                           diagnoses=[{"code": "M54.2", "description": "Cervicalgia"}]), PAGES, {2, 3}, CTX)
    assert r.needs_review and r.confidence == "low"
    assert "date 2025-03-15 not found" in r.review_note
    assert "removed ICD code M54.2" in r.review_note and r.diagnoses == [{"code": "", "description": "Cervicalgia"}]
    assert "800" in r.review_note


def test_pages_not_given_and_another_patients_page_are_caught():
    r = check_entry(_entry(pages=[2, 40]), PAGES, {2, 3}, CTX)
    assert r.pages == [2] and "cited a page it was not given" in r.review_note
    hale = check_entry(_entry(date="2025-04-08", provider="Mark Chen, PT, DPT", facility="", complaints="Knee pain 3/10",
                              diagnoses=[], treatment="", work_status="", pages=[9]), PAGES, {9}, CTX)
    assert hale.needs_review and "another patient" in hale.review_note
    jr = CheckContext(patient="Dana Whitfield Jr.", doi=date(2025, 3, 14), today=TODAY)
    assert not check_entry(_entry(), PAGES, {2, 3}, jr).needs_review


def test_records_before_the_injury_are_marked():
    pages = {5: "Patient: Dana Whitfield\nDate of visit: June 10, 2024\nProvider: Thomas Keller, DO\nLow back pain"}
    r = check_entry(ExtractedEntry(date="2024-06-10", provider="Thomas Keller, DO", visit_type="Office visit",
                                   complaints="Low back pain", pages=[5], confidence="high"), pages, {5}, CTX)
    assert r.pre_incident and not r.needs_review


def test_bill_amounts_and_codes_must_be_on_the_page():
    ok = check_bill(ExtractedBill(date="2025-03-14", provider="Maple Falls General Hospital", code="99284",
                                  description="ED visit", charge="$1,450.00", pages=[3]), PAGES, {3}, CTX)
    assert not ok.needs_review and ok.charge == 1450.0
    bad = check_bill(ExtractedBill(date="2025-03-14", provider="Maple Falls General Hospital", code="99285",
                                   charge=1540.0, pages=[3]), PAGES, {3}, CTX)
    assert bad.needs_review and bad.code == "" and "1,540.00 not found" in bad.review_note


# --- merging, gaps, bills -----------------------------------------------------------------------------------------

def _e(d, provider, pages, **kw):
    return Entry(date=d, provider=provider, pages=pages, confidence="high", needs_review=False, **kw)


def test_split_and_duplicate_records_merge_but_other_visits_do_not():
    page_file = {1: "er.pdf", 2: "er.pdf", 3: "er.pdf", 30: "insurer.pdf", 31: "insurer.pdf"}
    entries = [_e("2025-03-14", "Priya Raman, MD", [1], complaints="Neck pain 7/10"),
               _e("2025-03-14", "Dr. Raman", [2], treatment="Ibuprofen 600 mg"),
               _e("2025-03-14", "Laura Benson, MD", [3], visit_type="CT cervical spine"),
               _e("2025-03-14", "Priya Raman, MD", [30, 31], complaints="Neck pain 7/10")]
    merged = merge_entries(entries, page_file)
    assert len(merged) == 2
    er = next(e for e in merged if "Raman" in e.provider)
    assert er.pages == [1, 2] and er.duplicate_pages == [30, 31]
    assert er.complaints == "Neck pain 7/10" and er.treatment == "Ibuprofen 600 mg"


def test_bill_copies_from_another_file_are_dropped_but_units_stay():
    page_file = {5: "er.pdf", 6: "er.pdf", 40: "copy.pdf"}
    line = dict(date="2025-03-14", provider="MFGH", code="97110", charge=95.0)
    bills = [Bill(pages=[5], **line), Bill(pages=[5], **line), Bill(pages=[40], **line), Bill(pages=[6], **line)]
    out = merge_bills(bills, page_file)
    assert len(out) == 3  # the copy from copy.pdf is dropped
    assert out[2].needs_review and "duplicate or a second unit" in out[2].review_note


def test_gaps_skip_records_before_the_injury():
    entries = [_e("2024-06-10", "Keller", [1], pre_incident=True), _e("2025-03-14", "Raman", [2]),
               _e("2025-04-22", "Ortiz", [3]), _e("2025-06-18", "Okafor", [4])]
    gaps = find_gaps(entries, 30)
    assert [(g.start, g.end, g.days) for g in gaps] == [("2025-03-14", "2025-04-22", 39), ("2025-04-22", "2025-06-18", 57)]
    assert page_range([3, 4, 5, 9, 11, 12]) == "3-5, 9, 11-12"


# --- records and batching -----------------------------------------------------------------------------------------

def test_batches_stay_within_a_file_and_skip_scanned_pages():
    pages = [Page(n, "a.pdf" if n <= 7 else "b.pdf", n, "text " * 20) for n in range(1, 10)]
    pages[2].text = ""  # a scan
    batches = make_batches(pages, Options(batch_pages=3))
    assert [[p.n for p in b] for b in batches] == [[1, 2, 4], [5, 6, 7], [8, 9]]


def test_scanned_pages_are_listed_for_ocr_in_one_review_row(tmp_path):
    from medchron.pdfgen import text_pdf

    (tmp_path / "scan.pdf").write_bytes(text_pdf([["x"], [], [], ["Patient: Dana Whitfield " * 3]]))
    case = read_case(tmp_path)
    assert [p.needs_ocr for p in case.pages] == [True, True, True, False]
    c = build([], [], case, doi=None)
    rows = [r for r in c.review if r["kind"] == "page"]
    assert rows[0]["pages"] == "1-3" and "--ocr" in rows[0]["note"]
    assert rows[1]["pages"] == "4" and "check for anything missed" in rows[1]["note"]


def test_ocr_runs_ocrmypdf_and_reports_failure(tmp_path):
    src = tmp_path / "scan.pdf"
    src.write_bytes(b"%PDF-1.4")
    calls = []

    def fake_run(cmd, **kw):
        calls.append(cmd)
        Path(cmd[-1]).write_bytes(b"%PDF-1.4 ocr")
        return SimpleNamespace(returncode=0, stderr="")

    out = ocr_pdf(src, tmp_path / "ocr", run=fake_run)
    assert out.read_bytes().endswith(b"ocr") and calls[0][:2] == ["ocrmypdf", "--skip-text"]
    ocr_pdf(src, tmp_path / "ocr", run=fake_run)
    assert len(calls) == 1  # reused while the source is unchanged
    with pytest.raises(RuntimeError, match="ocrmypdf failed"):
        ocr_pdf(src, tmp_path / "ocr2", run=lambda cmd, **kw: SimpleNamespace(returncode=2, stderr="bad"))


# --- the synthetic case ---------------------------------------------------------------------------------------------

def test_committed_example_matches_the_generator(tmp_path):
    truth = make_sample(tmp_path)
    assert truth == json.loads((EXAMPLE / "truth.json").read_text())
    for pdf in (tmp_path / "records").glob("*.pdf"):
        assert pdf.read_bytes() == (EXAMPLE / "records" / pdf.name).read_bytes(), pdf.name


def _oracle(truth: dict, texts: dict[int, str], split_pages: bool = True):
    """A perfect stand-in model: answers from the answer key and reads bill lines off the pages."""

    def respond(kwargs):
        prompt = user_prompt(kwargs)
        shown = [int(n) for n in re.findall(r"^### Page (\d+)(?! \(context)", prompt, re.M)]
        entries, bills = [], []
        for t in truth["encounters"]:
            for copy in (t["pages"], t.get("duplicate_pages", [])):
                here = [p for p in copy if p in shown]
                if not here:
                    continue
                e = {"date": t["date"], "provider": t["provider"], "facility": t["facility"],
                     "visit_type": t["visit_type"], "confidence": "high"}
                for part in ([[p] for p in here] if split_pages else [here]):  # split visits must merge back
                    on_part = " ".join(texts[p] for p in part)
                    entries.append({**e, "pages": part, "diagnoses": [{"code": c, "description": ""}
                                                                      for c in t["icd"] if c in on_part]})
        for n in shown:
            for m in re.finditer(r"^(.+?) \| (\w+) \| ([^|]+) \| \d+ \| \$([\d,.]+)$", texts[n], re.M):
                bills.append({"date": parse_date(m.group(1)).isoformat(), "provider": texts[n].splitlines()[0],
                              "code": m.group(2),
                              "description": m.group(3).strip(), "charge": m.group(4), "pages": [n]})
        return json.dumps({"entries": entries, "bills": bills,
                           "no_content_pages": [p for p in truth["no_content_pages"] if p in shown],
                           "other_patient_pages": [p for p in truth["other_patient_pages"] if p in shown]})
    return respond


def _run_sample(tmp_path, client_responder=None, batch_pages=5):
    truth = make_sample(tmp_path / "case")
    case = read_case(tmp_path / "case" / "records")
    texts = {p.n: p.text for p in case.pages}
    client = FakeClient(client_responder or _oracle(truth, texts))
    llm = LLM(client, "test-model", rpm=0, sleep=no_sleep)
    opts = Options(patient=truth["patient"], doi=date.fromisoformat(truth["doi"]), batch_pages=batch_pages)
    run = asyncio.run(extract_all(case, opts, llm, tmp_path / "cache", progress=lambda _: None))
    return truth, case, run, llm, client


def test_end_to_end_on_the_synthetic_case_scores_full_marks(tmp_path):
    truth, case, run, llm, client = _run_sample(tmp_path)
    assert "Dana Whitfield" in client.calls[0]["messages"][0]["content"]
    c = build(run.entries, run.bills, case, doi=date.fromisoformat(truth["doi"]), no_content=run.no_content,
              other_patient=run.other_patient, not_processed=run.not_processed)
    from medchron.output import write_json

    write_json(tmp_path / "chronology.json", c)
    s = score(json.loads((tmp_path / "chronology.json").read_text()), truth)
    assert (s["found"], s["missed"], s["extra"]) == (16, [], [])
    assert s["codes_found"] == s["codes"] and s["wrong_codes"] == 0 and s["right_page"] == 16
    assert s["pre_incident_ok"] == 16 and s["duplicates_merged"] == 1
    assert s["gaps_found"] == 1 and s["extra_gaps"] == []
    assert s["bill_lines_found"] == s["bill_lines"] and s["bill_lines_extra"] == 0
    assert c.total_billed == truth["total_billed"] and c.days_to_first_treatment == 0
    assert s["flagged"] == 0
    kinds = [r["kind"] for r in c.review]
    assert "entry" not in kinds and "spot check" in kinds
    pages = [r for r in c.review if r["kind"] == "page"]  # every page is used or empty, except the misfiled one
    assert [(r["pages"], r["note"]) for r in pages] == [("15", "the model says this page is about another patient")]

    # a second run comes from the cache
    client2 = FakeClient(lambda kw: (_ for _ in ()).throw(AssertionError("should come from cache")))
    llm2 = LLM(client2, "test-model", rpm=0, sleep=no_sleep)
    opts = Options(patient=truth["patient"], doi=date.fromisoformat(truth["doi"]))
    run2 = asyncio.run(extract_all(case, opts, llm2, tmp_path / "cache", progress=lambda _: None))
    assert run2.batches_from_cache == run.batches and llm2.usage.requests == 0


def test_outputs_word_excel_and_bookmarked_pdf(tmp_path):
    import docx
    from openpyxl import load_workbook
    from pypdf import PdfReader

    from medchron.output import bookmarks, write_docx, write_xlsx

    truth, case, run, _, _ = _run_sample(tmp_path)
    run.entries[0].needs_review, run.entries[0].review_note = True, "date 2025-03-15 not found on the cited pages"
    doi = date.fromisoformat(truth["doi"])
    c = build(run.entries, run.bills, case, doi=doi, no_content=run.no_content, other_patient=run.other_patient)
    write_docx(tmp_path / "c.docx", c, case, patient="Dana Whitfield", doi=doi)
    d = docx.Document(tmp_path / "c.docx")
    text = "\n".join(p.text for p in d.paragraphs)
    assert "DRAFT: 1 yellow row" in text and "04/22/2025 to 06/18/2025 (57 days)" in text and "$11,928.00" in text
    chron = d.tables[1]
    assert chron.rows[1].cells[0].text.startswith("06/10/2024") and "Before injury" in chron.rows[1].cells[0].text
    assert "copy: 30-31" in chron.rows[2].cells[4].text

    write_xlsx(tmp_path / "c.xlsx", c, case)
    wb = load_workbook(tmp_path / "c.xlsx")
    assert wb.sheetnames == ["Chronology", "Treatment gaps", "Bills", "Bill totals", "Review", "Pages"]
    ws = wb["Chronology"]
    assert ws.cell(2, 1).value.date() == date(2024, 6, 10) and ws.cell(2, 2).value == "yes"
    assert wb["Bill totals"].cell(wb["Bill totals"].max_row, 3).value == 11928.0

    write_combined(case, tmp_path / "combined.pdf", bookmarks(c))
    r = PdfReader(tmp_path / "combined.pdf")
    assert len(r.pages) == 31 and "Page 2 of 31" in r.pages[1].extract_text()
    titles = [o.title for o in r.outline[1]]
    assert titles[0].startswith("06/10/2024 Office visit") and len(titles) == 16


def test_daily_limit_stops_and_marks_pages_not_processed(tmp_path):
    def capped(kwargs):
        req = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions")
        raise openai.RateLimitError("Rate limit exceeded: free-models-per-day", response=httpx.Response(429, request=req),
                                    body=None)

    truth, case, run, _, _ = _run_sample(tmp_path, capped)
    assert "daily request limit" in run.stopped_reason
    assert run.not_processed == {p.n for p in case.pages} and not run.entries
    c = build(run.entries, run.bills, case, doi=None, not_processed=run.not_processed)
    notes = {r["note"] for r in c.review}
    assert notes == {"not processed yet: re-run to continue"} and len(c.review) == 7  # one row per file


def test_the_model_over_reaching_is_caught_end_to_end(tmp_path):
    base = None

    def sloppy(kwargs):
        reply = json.loads(base(kwargs))
        for e in reply["entries"]:
            if e["visit_type"] == "MRI cervical spine":
                e["diagnoses"].append({"code": "M54.2", "description": "Cervicalgia"})  # not in the record
            if e["visit_type"] == "Orthopedic consultation":
                e["date"] = "2025-03-14"  # the date of injury from the history, not the visit date
        for b in reply["bills"]:
            if b["code"] == "72141":
                b["charge"] = 1785.0  # digits swapped
        return json.dumps(reply)

    truth = make_sample(tmp_path / "case")
    case = read_case(tmp_path / "case" / "records")
    base = _oracle(truth, {p.n: p.text for p in case.pages}, split_pages=False)
    llm = LLM(FakeClient(sloppy), "m", rpm=0, sleep=no_sleep)
    run = asyncio.run(extract_all(case, Options(patient=truth["patient"], doi=date(2025, 3, 14)), llm,
                                  tmp_path / "cache", progress=lambda _: None))
    flagged = {(e.visit_type, e.review_note) for e in run.entries if e.needs_review}
    assert any(v == "MRI cervical spine" and "removed ICD code M54.2" in n for v, n in flagged)
    consult = next(e for e in run.entries if e.visit_type == "Orthopedic consultation")
    assert consult.needs_review and "date of service as 2025-04-15, not 2025-03-14" in consult.review_note
    mri_bill = next(b for b in run.bills if b.code == "72141")
    assert mri_bill.needs_review and "1,785.00 not found" in mri_bill.review_note


# --- command line -----------------------------------------------------------------------------------------------

def test_real_records_need_a_baa_and_a_paid_model(monkeypatch):
    from medchron.__main__ import Settings, privacy_problem

    monkeypatch.delenv("MEDCHRON_BAA", raising=False)
    assert "Free models" in privacy_problem(Settings(model="nvidia/nemotron-3-super-120b-a12b:free"), False)
    assert privacy_problem(Settings(model="nvidia/nemotron-3-super-120b-a12b:free"), True) == ""
    assert "business associate agreement" in privacy_problem(Settings(model="anthropic/claude-sonnet"), False)
    monkeypatch.setenv("MEDCHRON_BAA", "yes")
    assert privacy_problem(Settings(model="anthropic/claude-sonnet"), False) == ""
    assert "Free models" in privacy_problem(Settings(model="openrouter/free"), False)


def test_cli_run_and_score(tmp_path, monkeypatch, capsys):
    from medchron import __main__ as cli

    truth = make_sample(tmp_path / "case")
    case = read_case(tmp_path / "case" / "records")
    oracle = _oracle(truth, {p.n: p.text for p in case.pages})
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(cli.LLM, "from_settings",
                        classmethod(lambda cls, s: LLM(FakeClient(oracle), s.model, rpm=0, sleep=no_sleep)))
    with pytest.raises(SystemExit, match="Free models"):
        cli.main(["run", "--records", str(tmp_path / "case" / "records"), "--model", "x:free"])
    out = tmp_path / "out"
    cli.main(["run", "--records", str(tmp_path / "case" / "records"), "--patient", "Dana Whitfield",
              "--doi", "2025-03-14", "--out", str(out), "--synthetic", "--model", "x:free"])
    for name in ("chronology.docx", "chronology.xlsx", "review.csv", "chronology.json", "records_combined.pdf",
                 "run_summary.json"):
        assert (out / name).exists(), name
    stats = json.loads((out / "run_summary.json").read_text())
    assert stats["entries"] == 16 and stats["treatment_gaps"] == 1 and stats["total_billed"] == 11928.0
    capsys.readouterr()
    cli.main(["score", "--run", str(out), "--truth", str(tmp_path / "case" / "truth.json")])
    report = capsys.readouterr().out
    assert "Encounters found: 16 of 16 (100%)" in report and "Total billed: $11,928.00" in report


def test_town_names_do_not_join_different_practices():
    from medchron.chronology import provider_summary, same_person, same_place

    assert same_place("Maple Falls General Hospital", "Maple Falls General Hospital ED")
    assert not same_place("Maple Falls General Hospital", "Maple Falls Physical Therapy")
    assert same_person("Priya Raman, MD", "Dr. Raman") and same_person("Mark Chen, PT, DPT", "Chen, Mark")
    assert not same_person("Mark Chen", "Mark Jones")
    rows = provider_summary([_e("2025-03-14", "Raman", [1], facility="Maple Falls General Hospital"),
                             _e("2025-03-24", "Chen", [2], facility="Maple Falls Physical Therapy"),
                             _e("2025-03-27", "Chen", [3], facility="Maple Falls Physical Therapy")])
    assert [(r["provider"], r["visits"]) for r in rows] == [("Maple Falls General Hospital", 1),
                                                            ("Maple Falls Physical Therapy", 2)]
