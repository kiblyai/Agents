import asyncio
import csv
import json
import re
from pathlib import Path

import httpx
import openai
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from agentkit.llm import LLM
from secq.docs import load_kb
from secq.draft import Drafted, Evidence, check
from secq.pipeline import Options, draft_all, review_rows
from secq.search import BM25, tokens
from secq.sheet import Question, detect, open_workbook, read_questions, write_draft

from .fakes import FakeClient, no_sleep, user_prompt

EXAMPLE_KB = Path(__file__).resolve().parents[1] / "examples" / "secq" / "kb"


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)


def _questionnaire(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Security"
    ws.append(["Vendor questionnaire"])
    ws.append([])
    ws.append(["ID", "Control Question", "Vendor Response", "Comments"])
    ws.append([None, "Encryption", None, None])
    ws.cell(4, 2).font = Font(bold=True)
    ws.append(["1", "Is customer data encrypted at rest?", None, None])
    ws.append(["2", "Describe your backup retention", None, None])
    ws.append(["3", "Are you ISO 27001 certified?", None, None])
    ws.append(["4", "Do you have a SOC 2 Type II report?", "Yes", "Already answered"])
    wb.save(path)
    return path


# --- search -----------------------------------------------------------------------------------------------------

def test_synonyms_bring_questions_and_policies_together():
    assert "mfa" in tokens("Is multi-factor authentication enforced?")
    assert "mfa" in tokens("We require 2FA for staff")
    idx = BM25(["Backups run nightly.", "MFA is required for all staff via Okta.", "We use TLS 1.2 in transit."])
    assert idx.top("Is multi-factor authentication enforced for employees?", 1)[0][0] == 1
    assert idx.top("Are you ISO 27001 certified?", 3) == []


# --- documents --------------------------------------------------------------------------------------------------

def test_knowledge_base_reads_every_supported_type(tmp_path):
    import docx

    (tmp_path / "policy.md").write_text("# Encryption\nData is encrypted at rest with AES-256.\n\n# Access\nMFA for all staff.")
    d = docx.Document()
    d.add_heading("Backups", level=1)
    d.add_paragraph("Backups are kept for 35 days.")
    d.save(tmp_path / "bcp.docx")
    (tmp_path / "soc2.pdf").write_bytes(_minimal_pdf("Penetration testing happens every year."))
    (tmp_path / "past.csv").write_text("Question,Answer\nDo you use subprocessors?,Yes - listed on our site\n")
    (tmp_path / "logo.png").write_bytes(b"\x89PNG")

    kb = load_kb(tmp_path)
    sources = [p.source for p in kb.passages]
    assert "policy.md § Encryption" in sources and "policy.md § Access" in sources
    assert "bcp.docx § Backups" in sources
    assert any(s == "soc2.pdf p.1" for s in sources)
    assert "Penetration testing" in next(p.text for p in kb.passages if p.source == "soc2.pdf p.1")
    assert kb.library[0].question == "Do you use subprocessors?" and kb.library[0].answer == "Yes - listed on our site"
    assert any("logo.png" in s for s in kb.files_skipped)


def test_example_knowledge_base_loads():
    kb = load_kb(EXAMPLE_KB)
    assert len(kb.passages) >= 10 and len(kb.library) == 2 and not kb.files_skipped


# --- questionnaire layout ---------------------------------------------------------------------------------------

def test_detects_layout_sections_and_prefilled(tmp_path):
    path = _questionnaire(tmp_path / "q.xlsx")
    wb = open_workbook(path)
    lay = detect(wb["Security"])
    assert (lay.header_row, lay.q_col, lay.a_col, lay.c_col, lay.id_col) == (3, 2, 3, 4, 1)
    layouts, qs = read_questions(wb)
    assert [q.qid for q in qs] == ["1", "2", "3", "4"]
    assert qs[1].text == "Describe your backup retention" and qs[1].section == "Encryption"  # not a section title
    assert [q.prefilled for q in qs] == [False, False, False, True]


def test_csv_questionnaire_and_column_override(tmp_path):
    p = tmp_path / "q.csv"
    p.write_text("Item,Text,Reply\n1,Do you encrypt backups?,\n")
    layouts, qs = read_questions(open_workbook(p))  # "Reply" is not a known header word
    assert qs == []
    layouts, qs = read_questions(open_workbook(p), override={"q_col": 2, "a_col": 3})
    assert [q.text for q in qs] == ["Do you encrypt backups?"]


# --- checks on drafted answers ----------------------------------------------------------------------------------

Q = Question(key="S!5", sheet="S", row=5, qid="1", text="Is data encrypted at rest?")


def _ev():
    from secq.docs import Passage
    return Evidence(passages={"S1-1": Passage("P1", "policy.md § Encryption", "Data is encrypted at rest with AES-256.")})


def test_well_sourced_answer_is_not_flagged():
    r = check(Drafted(id="Q1", answer="Yes", explanation="Customer data is encrypted at rest with AES-256.",
                      sources=["S1-1"], confidence="high", needs_review=False), Q, _ev())
    assert not r.needs_review and r.source_labels == ["policy.md § Encryption"]


def test_unsupported_numbers_and_certifications_are_flagged():
    r = check(Drafted(id="Q1", answer="Yes", explanation="We use AES-512 and are ISO 27001 certified.",
                      sources=["S1-1"], confidence="high", needs_review=False), Q, _ev())
    assert r.needs_review and r.confidence == "low"
    assert "512" in r.review_note and "iso27001" in r.review_note


def test_invented_source_ids_and_empty_answers_are_flagged():
    r = check(Drafted(id="Q1", answer="Yes", explanation="Encrypted.", sources=["S9-9"], confidence="high",
                      needs_review=False), Q, _ev())
    assert r.needs_review and "cited a source it was not given" in r.review_note and "no valid source" in r.review_note
    r2 = check(Drafted(id="Q1", confidence="high", needs_review=False), Q, _ev())
    assert r2.needs_review and r2.review_note == "Not covered by the documents provided"


# --- end to end with a fake model -------------------------------------------------------------------------------

def _responder(kwargs):
    prompt = user_prompt(kwargs)
    answers = []
    for block in prompt.split("### ")[1:]:
        qid = block.split("\n", 1)[0].strip()
        sources = re.findall(r"^\[(S\d+-\d+)\] \(([^)]*)\) (.*)$", block, re.M)
        question = re.search(r"^Question: (.*)$", block, re.M).group(1)
        if "ISO" in question:  # the model over-claims; our checks must catch it
            answers.append({"id": qid, "answer": "Yes", "explanation": "We are ISO 27001 certified.",
                            "sources": [sources[0][0]] if sources else [], "confidence": "high", "needs_review": False})
        elif "backup" in question.lower():
            continue  # the model skips a question
        elif sources:
            answers.append({"id": qid, "answer": "Yes", "explanation": sources[0][2], "sources": [sources[0][0]],
                            "confidence": "high", "needs_review": False})
    return json.dumps({"answers": answers})


def test_end_to_end_draft_written_back_and_cached(tmp_path):
    kb_dir = tmp_path / "kb"
    kb_dir.mkdir()
    (kb_dir / "policy.md").write_text("# Encryption\nCustomer data is encrypted at rest with AES-256.\n\n"
                                      "# Backups\nBackups are kept for 35 days.")
    q_path = _questionnaire(tmp_path / "q.xlsx")
    wb = open_workbook(q_path)
    layouts, qs = read_questions(wb)
    todo = [q for q in qs if not q.prefilled]
    kb = load_kb(kb_dir)
    client = FakeClient(_responder)
    llm = LLM(client, "test-model", rpm=0, sleep=no_sleep)
    run = asyncio.run(draft_all(todo, kb, Options(company="Acme"), llm, tmp_path / "cache", progress=lambda _: None))
    by_q = {r.qid: r for r in run.results}
    assert not by_q["1"].needs_review and by_q["1"].answer == "Yes"
    assert by_q["2"].needs_review and "skipped" in by_q["2"].review_note
    assert by_q["3"].needs_review and "iso27001" in by_q["3"].review_note
    assert "Acme" in client.calls[0]["messages"][0]["content"]

    out = tmp_path / "draft.xlsx"
    write_draft(q_path, out, layouts, run.results)
    ws = load_workbook(out)["Security"]
    assert ws.cell(5, 3).value == "Yes" and "AES-256" in ws.cell(5, 4).value
    assert ws.cell(5, 3).fill.start_color.rgb.endswith("FFF2CC") is False  # confident answer not highlighted
    assert ws.cell(7, 3).fill.start_color.rgb.endswith("FFF2CC")  # ISO over-claim highlighted
    assert ws.cell(3, 5).value == "Draft sources" and ws.cell(5, 5).value == "policy.md § Encryption"
    assert ws.cell(8, 3).value == "Yes" and ws.cell(8, 4).value == "Already answered"  # untouched

    rows = review_rows(run.results)
    assert rows[0]["needs_review"] == "yes" and rows[-1]["needs_review"] == ""

    client2 = FakeClient(lambda kw: (_ for _ in ()).throw(AssertionError("should come from cache")))
    llm2 = LLM(client2, "test-model", rpm=0, sleep=no_sleep)
    run2 = asyncio.run(draft_all(todo, kb, Options(company="Acme"), llm2, tmp_path / "cache", progress=lambda _: None))
    assert run2.batches_from_cache == 1 and llm2.usage.requests == 0


def test_daily_limit_stops_and_marks_the_rest(tmp_path):
    def capped(kwargs):
        req = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions")
        raise openai.RateLimitError("Rate limit exceeded: free-models-per-day", response=httpx.Response(429, request=req),
                                    body=None)

    kb = load_kb(EXAMPLE_KB)
    qs = [Question(key=f"S!{i}", sheet="S", row=i, qid=str(i), text=f"Is data encrypted at rest? ({i})") for i in range(1, 12)]
    llm = LLM(FakeClient(capped), "m", rpm=0, sleep=no_sleep)
    run = asyncio.run(draft_all(qs, kb, Options(company="Acme", batch_size=4), llm, tmp_path / "c", progress=lambda _: None))
    assert "daily request limit" in run.stopped_reason
    assert len(run.results) == 11 and all("not processed yet" in r.review_note for r in run.results)


def test_yes_must_be_backed_for_the_questions_own_claims_and_topic():
    from secq.docs import Passage

    soc = Evidence(passages={"S1-1": Passage("P1", "overview.md § Compliance", "We hold a SOC 2 Type II report.")})
    q_ins = Question(key="S!9", sheet="S", row=9, qid="4.3", text="Do you carry cyber insurance of at least $5M?")
    r = check(Drafted(id="Q1", answer="Yes", explanation="We hold a SOC 2 Type II report.", sources=["S1-1"],
                      confidence="high", needs_review=False), q_ins, soc)
    assert r.needs_review and "5" in r.review_note and "not seem to be about this question" in r.review_note

    ir = Evidence(passages={"S1-1": Passage("P2", "ir.md § Testing", "The incident response plan is tested once a year.")})
    q_247 = Question(key="S!10", sheet="S", row=10, qid="3.2", text="Do you have a 24/7 incident response capability?")
    r2 = check(Drafted(id="Q1", answer="Yes", explanation="The incident response plan is tested once a year.",
                       sources=["S1-1"], confidence="high", needs_review=False), q_247, ir)
    assert r2.needs_review and "24" in r2.review_note

    # a good answer to a short question is not over-flagged
    bk = Evidence(passages={"S1-1": Passage("P3", "dr.md § Backups", "Backups are taken daily, encrypted, and kept for 35 days.")})
    q_bk = Question(key="S!11", sheet="S", row=11, qid="1.3", text="How long are backups retained?")
    r3 = check(Drafted(id="Q1", answer="35 days", explanation="Backups are kept for 35 days.", sources=["S1-1"],
                       confidence="high", needs_review=False), q_bk, bk)
    assert not r3.needs_review


def test_section_title_does_not_outrank_the_question():
    from secq.pipeline import gather_evidence

    kb = load_kb(EXAMPLE_KB)
    q = Question(key="S!11", sheet="S", row=11, qid="2.2", text="Do you perform background checks on employees?",
                 section="Access Control")
    ev = gather_evidence([q], kb, Options(company="x"))[0]
    assert next(iter(ev.passages.values())).source == "security-overview.md § People"


def test_hyphenated_phrases_and_247_find_the_right_passage():
    from secq.pipeline import gather_evidence

    kb = load_kb(EXAMPLE_KB)
    q = Question(key="S!14", sheet="S", row=14, qid="3.2", text="Do you have a 24/7 incident response capability?",
                 section="Incident Response")
    ev = gather_evidence([q], kb, Options(company="x"))[0]
    assert next(iter(ev.passages.values())).source == "incident-response.md § Coverage"


def test_the_breach_notice_case_gets_no_unrelated_past_answer():
    # stand-in run on the example: "notify customers of a data breach" was offered the subprocessors past answer
    # ("customers are notified 30 days before changes") because both mention customer data, and a "Yes" built on
    # it passed every check
    from secq.pipeline import gather_evidence

    kb = load_kb(EXAMPLE_KB)
    qs = [Question(key="S!13", sheet="S", row=13, qid="3.1",
                   text="How quickly do you notify customers of a confirmed data breach?"),
          Question(key="S!6", sheet="S", row=6, qid="1.1", text="Is customer data encrypted at rest?"),
          Question(key="S!16", sheet="S", row=16, qid="4.1", text="Do you have a current SOC 2 Type II report?")]
    breach, encryption, soc2 = gather_evidence(qs, kb, Options(company="x"))
    assert not breach.library and not encryption.library
    past = list(soc2.library.values())
    assert [a.question for a in past] == ["Do you have a SOC 2 Type II report?"]
    assert past[0].answer.startswith("Yes. Our SOC 2 Type II report covers")


def test_numbers_must_match_whole_not_as_part_of_another_number():
    from secq.docs import Passage

    bk = Evidence(passages={"S1-1": Passage("P3", "dr.md § Backups", "Backups are encrypted with AES-256 and kept for 35 days.")})
    q = Question(key="S!12", sheet="S", row=12, qid="1.4", text="Are backups retained for at least 3 months?")
    r = check(Drafted(id="Q1", answer="Yes", explanation="Backups are kept for 35 days.", sources=["S1-1"],
                      confidence="high", needs_review=False), q, bk)
    assert r.needs_review and "mentions 3 " in r.review_note  # "3" is not backed by "35 days"

    ok = Question(key="S!13", sheet="S", row=13, qid="1.5", text="What encryption protects backups?")
    r2 = check(Drafted(id="Q1", answer="AES-256", explanation="Backups are encrypted with AES-256.", sources=["S1-1"],
                       confidence="high", needs_review=False), ok, bk)
    assert not r2.needs_review
