import asyncio
import csv
import json
import re
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path

import httpx
import openai
import pytest

from agentkit.llm import LLM
from freightq.broker import BrokerConfig, load_broker
from freightq.extract import (Email, EmailReading, ExtractedShipment, Shipment, amounts, check_reading,
                              check_shipment, find_dates, hazmat_mentions, zip_state)
from freightq.inbox import read_inbox, split_thread
from freightq.output import LOAD_COLUMNS, REVIEW_COLUMNS, load_rows, review_rows, write_csv, write_json, write_replies
from freightq.pipeline import read_all
from freightq.pricing import LaneLoad, load_lanes, price
from freightq.sample import make_sample, score

from .fakes import FakeClient, no_sleep, user_prompt

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "freightq"
CENTRAL = timezone(timedelta(hours=-5))
MON = datetime(2026, 9, 21, 8, 12, tzinfo=CENTRAL)  # a Monday
CFG = BrokerConfig()


def _email(body: str, sent: datetime | None = MON, subject: str = "Rate request", **kw) -> Email:
    latest, earlier = split_thread(body)
    return Email(file="test.eml", sender="maria@wwi.example", sender_name="Maria Lopez", subject=subject, sent=sent,
                 latest=latest, earlier=earlier, **kw)


DAL_ATL = ("Hi Coho team,\n\nPick up: Dallas, TX 75207\nDeliver: Atlanta, GA 30318\n"
           "Pickup Thursday 9/24, 8am-2pm. Deliver by Monday 9/28.\n"
           "53' dry van, 22 pallets of packaged housewares, 38,500 lbs.\n")


def _x(**kw) -> ExtractedShipment:
    base = dict(origin={"city": "Dallas", "state": "TX", "zip": "75207"},
                destination={"city": "Atlanta", "state": "GA", "zip": "30318"}, pickup_date="2026-09-24",
                pickup_time="8am-2pm", delivery_date="2026-09-28", equipment="van", weight_lbs=38500,
                commodity="packaged housewares", pieces="22 pallets")
    base.update(kw)
    return ExtractedShipment.model_validate(base)


def _check(x: ExtractedShipment, body: str = DAL_ATL, kind: str = "quote_request", lanes: int = 1) -> Shipment:
    e = _email(body)
    return check_shipment(x, e, kind, "Wide World Importers", 1, CFG, e.day, hazmat_mentions(e.all_text()), lanes)


def codes(s) -> set[str]:
    return {f.code for f in s.flags}


# --- reading emails -----------------------------------------------------------------------------------------------

def test_threads_split_at_the_first_quoted_message():
    latest, earlier = split_thread("Make it Tempe.\n\nOn Tue, Sep 22, 2026 at 4:10 PM Jin Park <jin@x.example>\n"
                                   "wrote:\n> Salinas to Phoenix\n")
    assert latest == "Make it Tempe." and earlier.startswith("On Tue") and "> Salinas" in earlier
    outlook = "New date: Friday.\n\nFrom: Jin\nSent: Monday\nTo: quotes\nSubject: x\n\nold"
    assert split_thread(outlook)[0] == "New date: Friday."
    assert split_thread("ok\n-----Original Message-----\nold")[1].startswith("-----Original")
    assert split_thread("no history here") == ("no history here", "")


def test_inbox_reads_eml_with_pdf_html_and_pasted_text(tmp_path):
    from agentkit.pdfgen import text_pdf

    msg = EmailMessage()
    msg["From"], msg["Subject"], msg["Date"] = "Chris <c@aw.example>", "Tender 88213", "Tue, 22 Sep 2026 15:20:00 -0500"
    msg.set_content("<p>See the <b>attached</b> tender.</p>", subtype="html")
    msg.add_attachment(text_pdf([["Load #: 88213", "Joliet, IL 60436"]]), maintype="application", subtype="pdf",
                       filename="tender.pdf")
    msg.add_attachment(b"\x00\x01", maintype="image", subtype="png", filename="logo.png")
    (tmp_path / "a.eml").write_bytes(msg.as_bytes())
    (tmp_path / "b.txt").write_text("From: Sam <sam@ash.example>\nSubject: van out of Phoenix\n"
                                    "Date: Wed, 23 Sep 2026 16:05:00 -0700\n\nWhat for a van out of Phoenix?\n")
    (tmp_path / "notes.docx").write_bytes(b"x")
    emails, skipped = read_inbox(tmp_path)
    a, b = emails
    assert a.sender == "c@aw.example" and a.day == date(2026, 9, 22) and a.latest == "See the attached tender."
    assert a.attachments[0].name == "tender.pdf" and "Load #: 88213" in a.attachments[0].text
    assert a.skipped == ["logo.png"]
    assert b.sender == "sam@ash.example" and b.subject == "van out of Phoenix" and b.day == date(2026, 9, 23)
    assert b.latest == "What for a van out of Phoenix?" and skipped == ["notes.docx (not .eml or .txt)"]


def test_dates_are_read_from_the_day_the_email_was_sent():
    ref = date(2026, 9, 21)  # Monday
    assert date(2026, 9, 24) in find_dates("pickup Thursday", ref)
    assert find_dates("pick up tomorrow", ref) == {date(2026, 9, 22)}
    assert find_dates("next Tuesday", ref) == {date(2026, 9, 22), date(2026, 9, 29)}  # ambiguous: both
    assert find_dates("Monday", ref) == {ref, date(2026, 9, 28)}
    assert find_dates("9/24, Sep 25th, 26 Sep, 2026-09-27, 09/28/2026 and 9/29/26", ref) == {
        date(2026, 9, d) for d in range(24, 30)}
    assert find_dates("24/7 dock, $1/2 mile, 3.5/10", ref) == set()
    assert find_dates("ship 1/4", date(2026, 12, 28)) == {date(2027, 1, 4)}  # next year
    assert find_dates("delivered 12/30", date(2027, 1, 2)) == {date(2026, 12, 30)}  # last year


def test_amounts_hazmat_words_and_zip_states():
    assert {38500.0, 42000.0} <= amounts("38,500 lbs and 42k and 21 tons")
    assert hazmat_mentions("boxed toys, no hazmat, non-hazardous, class 55 freight") == []
    assert hazmat_mentions("paint, UN1263, Class 3, placards required") == ["UN1263", "Class 3", "placards"]
    assert (zip_state("75207"), zip_state("30318"), zip_state("85281"), zip_state("20166"), zip_state("0")) == \
        ("TX", "GA", "AZ", "VA", "")


# --- checks on one lane -------------------------------------------------------------------------------------------

def test_a_well_supported_lane_passes_with_nothing_missing():
    s = _check(_x())
    assert s.flags == [] and s.missing == [] and s.equipment == "van" and s.weight_lbs == 38500
    assert s.origin == {"city": "Dallas", "state": "TX", "zip": "75207"} and s.pickup_date == "2026-09-24"
    both = _check(_x(origin={"city": "Dallas, TX", "state": "", "zip": "75207"}))  # state left in the city field
    assert both.origin == s.origin and both.flags == []


def test_values_not_in_the_email_are_removed_and_asked_for():
    s = _check(_x(origin={"city": "Dallas", "state": "TX", "zip": "75201"}, weight_lbs=42000, pickup_date="2026-09-25",
                  references=["PO 5512"], rate=2100))
    notes = " | ".join(f.text for f in s.flags)
    assert s.origin["zip"] == "" and "removed origin ZIP 75201" in notes
    assert s.weight_lbs is None and "removed weight 42,000 lbs" in notes
    assert s.pickup_date == "" and "removed pickup date 2026-09-25" in notes
    assert s.references == [] and s.rate is None
    tender = _check(_x(references=["8821", "PO 4471-A"]), DAL_ATL + "Load 88213, PO 4471-A\n", kind="tender")
    assert tender.references == ["PO 4471-A"] and "removed reference '8821'" in tender.flags[0].text
    assert s.missing == ["pickup_date", "weight"]


def test_equipment_the_email_does_not_state_is_removed():
    s = _check(_x(equipment="reefer"))
    assert s.equipment == "" and "equipment" in s.missing and "removed equipment 'reefer'" in s.flags[0].text
    assert "temperature" not in s.missing


def test_reefer_needs_a_temperature_and_zips_must_match_states():
    body = "Reefer from Fresno, CA 93725 to Denver, TX 80216 tomorrow, 40,000 lbs grapes"
    s = _check(_x(origin={"city": "Fresno", "state": "CA", "zip": "93725"},
                  destination={"city": "Denver", "state": "TX", "zip": "80216"}, pickup_date="2026-09-22",
                  delivery_date="", equipment="reefer", weight_lbs=40000, commodity="grapes", pieces=""), body)
    assert s.missing == ["temperature"]
    assert [f.text for f in s.flags if f.code == "zip_state"] == ["destination ZIP 80216 is in CO, not TX"]


def test_swapped_origin_and_destination_are_caught():
    swapped = _check(_x(origin={"city": "Atlanta", "state": "GA", "zip": "30318"},
                        destination={"city": "Dallas", "state": "TX", "zip": "75207"}))
    assert "swap" in codes(swapped)
    same_line = "Need a van from Dallas, TX 75207 to Atlanta, GA 30318 Thursday 9/24, 38,500 lbs housewares"
    ok = _check(_x(delivery_date="", pickup_time="", pieces=""), same_line)
    assert "swap" not in codes(ok)
    flipped = _check(_x(origin={"city": "Atlanta", "state": "GA", "zip": "30318"},
                        destination={"city": "Dallas", "state": "TX", "zip": "75207"}, delivery_date="",
                        pickup_time="", pieces=""), same_line)
    assert "swap" in codes(flipped)
    reverse = "Going to Atlanta, GA 30318 from Dallas, TX 75207 on Thursday 9/24, 38,500 lbs housewares, van"
    assert "swap" not in codes(_check(_x(delivery_date="", pickup_time="", pieces=""), reverse))


def test_hazmat_follows_the_email_not_the_model():
    body = DAL_ATL + "Paint, UN1263, Class 3.\n"
    missed = _check(_x(hazmat=False), body)
    assert missed.hazmat and "hazmat" in missed.charged and "model missed it" in missed.flags[0].text
    invented = _check(_x(hazmat=True))
    assert not invented.hazmat and "removed hazmat" in invented.flags[0].text
    two_lanes = _check(_x(hazmat=False), body, lanes=2)
    assert not two_lanes.hazmat and "check whether this lane is" in two_lanes.flags[0].text
    assert _check(_x(), DAL_ATL + "No hazmat.\n").flags == []


def test_the_coho_team_greeting_is_not_a_team_driver_request():
    s = _check(_x(requirements=["team drivers", "tarps"]))
    assert s.requirements == [] and s.charged == []
    assert {f.text for f in s.flags} == {"removed requirement 'team drivers' (not in the email)",
                                         "removed requirement 'tarps' (not in the email)"}
    ok = _check(_x(requirements=["team drivers"]), DAL_ATL + "Needs team drivers.\n")
    assert ok.charged == ["team"] and ok.flags == []


def test_overweight_small_shipments_and_tenders_without_numbers_are_flagged():
    heavy = _check(_x(weight_lbs=47500), DAL_ATL.replace("38,500", "47,500"))
    assert codes(heavy) == {"overweight"}
    small_body = DAL_ATL.replace("38,500", "1,800").replace("22 pallets", "3 pallets")
    small = _check(_x(weight_lbs=1800, pieces="3 pallets"), small_body)
    assert codes(small) == {"small_shipment"}
    tender = _check(_x(), kind="tender")
    assert codes(tender) == {"reference"}


def test_email_level_checks():
    e = _email(DAL_ATL, sent=None, skipped=["scan.tif"])
    r = check_reading(EmailReading.model_validate({"kind": "quote", "shipments": [_x().model_dump()]}), e, CFG,
                      today=date(2026, 9, 21))
    assert r.kind == "quote_request" and len(r.shipments) == 1
    assert [f.code for f in r.flags] == ["date", "attachment"]
    carrier = check_reading(EmailReading(kind="other", shipments=[_x()]), _email("Empty van in Dallas"), CFG)
    assert carrier.shipments == [] and carrier.flags[0].code == "kind"


def test_messy_model_replies_are_tolerated():
    r = EmailReading.model_validate({"kind": "Load Tender", "shipments": [
        {"origin": "Joliet, IL 60436", "destination": {"city": "Nashville", "state": "TN"}, "weight_lbs": "29.75k",
         "rate": "$1,850.00", "hazmat": "no", "references": "88213, PO 4471-A", "truckloads": "2 loads",
         "extra_stops": ["Louisville, KY 40218"]}, "junk"]})
    s = r.shipments[0]
    assert r.kind == "tender" and len(r.shipments) == 1
    assert s.origin.zip == "60436" and s.extra_stops[0].city == "Louisville" and s.weight_lbs == 29750
    assert s.rate == 1850 and s.hazmat is False and s.references == ["88213", "PO 4471-A"] and s.truckloads == 2


# --- pricing ------------------------------------------------------------------------------------------------------

def _load(day, pay, rate=None, o=("Dallas", "TX", "75207"), d=("Atlanta", "GA", "30318"), eq="van", cust="Wide World"):
    return LaneLoad(1, date.fromisoformat(day), cust, dict(zip(("city", "state", "zip"), o)),
                    dict(zip(("city", "state", "zip"), d)), eq, None, pay, rate)


def _lane(**kw) -> Shipment:
    base = dict(kind="quote_request", customer="Wide World Importers", origin={"city": "Dallas", "state": "TX",
                "zip": "75207"}, destination={"city": "Atlanta", "state": "GA", "zip": "30318"}, equipment="van")
    base.update(kw)
    return Shipment(**base)


ASOF = date(2026, 9, 21)


def test_price_is_the_median_of_recent_matching_loads_plus_margin():
    loads = [_load("2026-09-15", 1875, 2250), _load("2026-09-02", 1950), _load("2026-08-20", 1800),
             _load("2026-08-05", 1900), _load("2026-07-21", 1850), _load("2026-06-10", 2300),
             _load("2026-01-15", 1500), _load("2026-09-10", 2400, eq="reefer"), _load("2026-09-30", 900)]
    s = _lane()
    q = price(s, loads, CFG, ASOF)
    assert q.basis == "same cities" and len(q.comparables) == 5 and q.buy == 1875
    assert q.linehaul == 2225 and q.total == 2225 and round(q.margin, 3) == 0.157 and s.flags == []


def test_minimum_margin_rounding_and_extra_charges():
    s = _lane(stops=[{"city": "Macon", "state": "GA", "zip": ""}] * 2, charged=["extra_stop", "tarps"])
    q = price(s, [_load("2026-09-15", 800), _load("2026-09-14", 810)], CFG, ASOF)
    assert q.buy == 805 and q.linehaul == 975  # 805 + $150 minimum beats 15%, rounded up to $25
    assert q.extras == [("extra_stop", 200.0), ("tarps", 125.0)] and q.total == 1300


def test_nearby_zips_then_states_then_nothing():
    near = _lane(origin={"city": "Memphis", "state": "TN", "zip": "38118"},
                 destination={"city": "Columbus", "state": "OH", "zip": "43207"})
    loads = [_load("2026-09-01", 1550, o=("Germantown", "TN", "38138"), d=("Obetz", "OH", "43207")),
             _load("2026-07-28", 1500, o=("Germantown", "TN", "38138"), d=("Obetz", "OH", "43207")),
             _load("2026-09-09", 900, o=("Memphis", "TN", "38118"), d=("St. Louis", "MO", "63147"))]
    q = price(near, loads, CFG, ASOF)
    assert q.basis == "nearby ZIPs" and q.linehaul == 1800 and codes(near) == {"nearby_zips"}
    kc = _lane(origin={"city": "Memphis", "state": "TN", "zip": "38118"},
               destination={"city": "Kansas City", "state": "MO", "zip": "64120"})
    q = price(kc, loads, CFG, ASOF)
    assert not q.priced and q.basis.startswith("same states") and codes(kc) == {"state_history_only"}
    none = _lane(destination={"city": "Boise", "state": "ID", "zip": "83716"})
    assert not price(none, loads, CFG, ASOF).priced and codes(none) == {"no_history"}


def test_pricing_warnings():
    old = _lane()
    price(old, [_load("2026-03-01", 1800), _load("2026-02-01", 1850)], CFG, ASOF)
    assert codes(old) == {"stale_history"}
    one = _lane()
    price(one, [_load("2026-09-01", 1800)], CFG, ASOF)
    assert codes(one) == {"few_loads"}
    wide = _lane()
    price(wide, [_load("2026-09-01", 1500), _load("2026-09-02", 2200)], CFG, ASOF)
    assert codes(wide) == {"wide_spread"}
    cust = _lane()
    price(cust, [_load("2026-09-01", 1800, 2500), _load("2026-09-02", 1800)], CFG, ASOF)
    assert codes(cust) == {"customer_rate"} and "paid $2,500" in cust.flags[0].text


def test_tender_rate_is_checked_against_carrier_pay():
    loads = [_load("2026-09-10", 1500), _load("2026-08-27", 1450), _load("2026-08-06", 1550)]
    good = _lane(kind="tender", rate=1850.0)
    q = price(good, loads, CFG, ASOF)
    assert q.total == 1850 and round(q.margin, 3) == 0.189 and good.flags == []
    thin = _lane(kind="tender", rate=1600.0)
    price(thin, loads, CFG, ASOF)
    assert codes(thin) == {"low_margin"} and "leaves $100" in thin.flags[0].text
    no_rate = _lane(kind="tender")
    assert price(no_rate, loads, CFG, ASOF).total == 1775 and codes(no_rate) == {"tender_rate"}


def test_lane_history_accepts_tms_column_names_and_skips_bad_rows(tmp_path):
    p = tmp_path / "lanes.csv"
    p.write_text("Pickup Date,Shipper,Pickup City,Pickup State,Pickup Zip,Delivery City,Delivery State,Delivery Zip,"
                 "Trailer Type,Carrier Rate,Revenue\n"
                 "09/15/2026,WWI,Dallas,Texas,75207,Atlanta,GA,30318,Dry Van,\"$1,875.00\",\"$2,250.00\"\n"
                 "09/14/2026,WWI,Dallas,TX,,Atlanta,GA,,Hotshot,900,\n"
                 "09/13/2026,WWI,Dallas,TX,,Atlanta,GA,,Van,,\n"
                 "09/12/2026,WWI,,,,Atlanta,GA,,Van,900,\n")
    loads, problems = load_lanes(p)
    assert len(loads) == 1 and loads[0].day == date(2026, 9, 15) and loads[0].origin["state"] == "TX"
    assert loads[0].carrier_pay == 1875 and loads[0].customer_rate == 2250
    assert problems == ["row 3: unknown equipment 'Hotshot'", "row 4: no carrier pay", "row 5: no origin or destination"]
    (tmp_path / "bad.csv").write_text("date,city\n")
    with pytest.raises(ValueError, match="needs columns"):
        load_lanes(tmp_path / "bad.csv")


def test_broker_settings(tmp_path):
    cfg = load_broker(EXAMPLE / "broker.toml")
    assert cfg.name == "Coho Freight Brokerage" and cfg.accessorials["tarps"] == 125 and cfg.max_weight_lbs["van"] == 45000
    (tmp_path / "b.toml").write_text("[broker]\ntarget_margin = 15\n")
    with pytest.raises(ValueError, match="share of the rate"):
        load_broker(tmp_path / "b.toml")
    (tmp_path / "c.toml").write_text("[broker]\nmargin = 0.15\n")
    with pytest.raises(ValueError, match="unknown"):
        load_broker(tmp_path / "c.toml")


# --- the synthetic inbox ------------------------------------------------------------------------------------------

def test_committed_example_matches_the_generator(tmp_path):
    truth = make_sample(tmp_path)
    assert truth == json.loads((EXAMPLE / "truth.json").read_text())
    for f in [*(tmp_path / "inbox").iterdir(), tmp_path / "lanes.csv", tmp_path / "broker.toml"]:
        rel = f.relative_to(tmp_path)
        assert f.read_bytes() == (EXAMPLE / rel).read_bytes(), rel


FIELDS = ("pickup_date", "pickup_time", "delivery_date", "delivery_time", "equipment", "weight_lbs", "commodity",
          "pieces", "temperature", "hazmat", "requirements", "references", "rate", "truckloads")


def _oracle(truth: dict):
    """A perfect stand-in model: answers from the answer key."""
    by_file = {e["file"]: e for e in truth["emails"]}

    def respond(kwargs):
        t = by_file[re.search(r"^Email file: (.+)$", user_prompt(kwargs), re.M).group(1)]
        ships = [{"origin": s["origin"], "destination": s["destination"], "extra_stops": s["stops"],
                  **{k: s[k] for k in FIELDS}} for s in t["shipments"]]
        return json.dumps({"kind": t["kind"], "customer": t["customer"], "contact": t["contact"], "shipments": ships})
    return respond


def _run(tmp_path, responder=None):
    truth = make_sample(tmp_path / "case")
    emails, _ = read_inbox(tmp_path / "case" / "inbox")
    cfg = load_broker(tmp_path / "case" / "broker.toml")
    lanes, problems = load_lanes(tmp_path / "case" / "lanes.csv")
    client = FakeClient(responder or _oracle(truth))
    llm = LLM(client, "test-model", rpm=0, sleep=no_sleep)
    run = asyncio.run(read_all(emails, cfg, lanes, llm, tmp_path / "cache", progress=lambda _: None))
    return truth, emails, cfg, lanes, problems, run, llm


def test_end_to_end_on_the_synthetic_inbox_scores_full_marks(tmp_path):
    truth, emails, cfg, lanes, problems, run, llm = _run(tmp_path)
    assert problems == ["row 43: unknown equipment 'Hotshot'"] and llm.usage.requests == 11
    write_replies(run, cfg, tmp_path / "replies")
    write_json(tmp_path / "results.json", run)
    s = score(json.loads((tmp_path / "results.json").read_text()), truth)
    assert (s["kind_right"], s["lanes_found"], s["missed"], s["extra_lanes"]) == (11, 12, [], 0)
    assert all(a == b for a, b in s["fields"].values()) and s["missing_right"] == 12
    assert (s["rate_right"], s["priced_expected"], s["wrong_rates"]) == (9, 9, [])
    assert (s["unpriced_right"], s["unpriced_expected"]) == (3, 3)
    assert (s["flags_found"], s["flags_expected"], s["extra_flags"]) == (6, 6, [])

    replies = {p.name: p.read_text() for p in (tmp_path / "replies").iterdir()}
    assert "09_carrier_capacity.txt" not in replies and len(replies) == 10
    first = replies["01_wwi_dallas_atlanta.txt"]
    assert first.startswith("To: maria.lopez@wideworldimporters.example\nSubject: Re: Rate request")
    assert "Rate: $2,225 all-in" in first and "Pickup Thu 09/24/2026 8am-2pm; deliver Mon 09/28/2026" in first
    assert "Rate: $2,150 all-in ($2,025 linehaul + $125 tarps)" in replies["03_fabrikam_flatbed.txt"]
    assert "Please confirm the temperature setting." in replies["02_contoso_fresno_denver.txt"]
    three = replies["04_tailspin_three_lanes.txt"]
    assert "Rate: $1,600" in three and "Rate: $1,800" in three and "[[PRICE BY HAND]]" in three
    assert "Confirming we'll cover this load at $1,850 all-in." in replies["05_adventure_tender.txt"]
    assert "$2,775 all-in per truck, 2 trucks" in replies["07_contoso_thread.txt"]
    assert "need the delivery city and state (or ZIP) and pickup date" in replies["11_alpine_phoenix.txt"]

    review = review_rows(run)
    assert review[0]["order"] == 1 and review[0]["check"].startswith(
        "tender: accept or decline; rate $1,850 vs recent carrier pay $1,500 (19% margin)")
    assert [r["email"] for r in review if r["order"] == 4] == ["09_carrier_capacity.eml"]
    assert not any(r["email"].startswith(("01", "02", "03")) for r in review)
    write_csv(tmp_path / "review.csv", review, REVIEW_COLUMNS)
    tenders = load_rows(run)
    assert len(tenders) == 1 and tenders[0]["load_number"] == "88213" and tenders[0]["stops"] == "Louisville, KY 40218"
    assert tenders[0]["customer_rate"] == "1850.00" and tenders[0]["est_carrier_pay"] == "1500.00"
    write_csv(tmp_path / "loads.csv", tenders, LOAD_COLUMNS)
    with open(tmp_path / "loads.csv") as f:
        assert next(csv.DictReader(f))["pickup_zip"] == "60436"

    # a second run comes from the cache
    client2 = FakeClient(lambda kw: (_ for _ in ()).throw(AssertionError("should come from cache")))
    llm2 = LLM(client2, "test-model", rpm=0, sleep=no_sleep)
    run2 = asyncio.run(read_all(emails, cfg, lanes, llm2, tmp_path / "cache", progress=lambda _: None))
    assert run2.from_cache == 11 and llm2.usage.requests == 0


def test_the_model_over_reaching_is_caught_end_to_end(tmp_path):
    base = None

    def sloppy(kwargs):
        reply = json.loads(base(kwargs))
        for s in reply["shipments"]:
            if s["destination"]["city"] == "Tempe":
                s["destination"] = {"city": "Phoenix", "state": "AZ", "zip": "85043"}  # the old message, not the newest
            if s["origin"]["city"] == "Houston" and s["destination"]["city"] == "Little Rock":
                s["hazmat"] = False  # missed
            if s["destination"]["city"] == "Atlanta":
                s["weight_lbs"] = 42000  # not in the email
                s["requirements"] = ["team drivers"]  # "Hi Coho team"
            if s["origin"]["city"] == "Phoenix":
                s["destination"] = {"city": "Boston", "state": "MA", "zip": ""}  # a guess
        return json.dumps(reply)

    truth = make_sample(tmp_path / "case")
    base = _oracle(truth)
    _, _, _, _, _, run, _ = _run(tmp_path, sloppy)
    by = {r.email.file: r for r in run.results}
    dal = by["01_wwi_dallas_atlanta.eml"].shipments[0]
    assert dal.weight_lbs is None and "weight" in dal.missing and dal.charged == []
    haz = by["08_proseware_hazmat.eml"].shipments[0]
    assert haz.hazmat and haz.quote.total == 1475 and "model missed it" in haz.flags[0].text
    phx = by["11_alpine_phoenix.txt"].shipments[0]
    assert "destination city 'Boston' is not in the email" in [f.text for f in phx.flags]
    write_json(tmp_path / "results.json", run)
    s = score(json.loads((tmp_path / "results.json").read_text()), truth)
    assert "07_contoso_thread.eml lane 1" in s["missed"]  # the scorer catches the wrong consignee


def test_daily_limit_stops_and_leaves_the_rest_for_the_next_run(tmp_path):
    calls = []

    def capped(kwargs):
        calls.append(1)
        if len(calls) > 2:
            req = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions")
            raise openai.RateLimitError("Rate limit exceeded: free-models-per-day",
                                        response=httpx.Response(429, request=req), body=None)
        return base(kwargs)

    truth = make_sample(tmp_path / "case")
    base = _oracle(truth)
    _, _, _, _, _, run, _ = _run(tmp_path, capped)
    assert "daily request limit" in run.stopped_reason
    assert [bool(r.error) for r in run.results] == [False, False] + [True] * 9
    rows = review_rows(run)
    assert len([r for r in rows if r["check"] == "not processed yet: re-run to continue"]) == 9


# --- command line -------------------------------------------------------------------------------------------------

def test_cli_run_score_and_price(tmp_path, monkeypatch, capsys):
    from freightq import __main__ as cli

    case = tmp_path / "case"
    cli.main(["sample", "--out", str(case)])
    truth = json.loads((case / "truth.json").read_text())
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(cli.LLM, "from_settings",
                        classmethod(lambda cls, s: LLM(FakeClient(_oracle(truth)), s.model, rpm=0, sleep=no_sleep)))
    out = tmp_path / "out"
    cli.main(["run", "--emails", str(case / "inbox"), "--lanes", str(case / "lanes.csv"), "--broker",
              str(case / "broker.toml"), "--out", str(out), "--model", "x:free"])
    printed = capsys.readouterr().out
    assert "WARNING: free models may log prompts" in printed and "9 of 12 lanes priced; 10 replies drafted" in printed
    for name in ("quotes.xlsx", "review.csv", "loads.csv", "results.json", "run_summary.json", "replies"):
        assert (out / name).exists(), name
    stats = json.loads((out / "run_summary.json").read_text())
    assert (stats["tenders"], stats["other"], stats["lanes"], stats["lanes_priced"]) == (1, 1, 12, 9)

    from openpyxl import load_workbook
    wb = load_workbook(out / "quotes.xlsx")
    assert wb.sheetnames == ["Shipments", "Comparable loads", "Emails", "Review"]
    ws = wb["Shipments"]
    assert ws.cell(2, 23).value == 2225 and ws.cell(2, 25).value == "same cities, 5 loads"

    cli.main(["score", "--run", str(out), "--truth", str(case / "truth.json")])
    report = capsys.readouterr().out
    assert "Lanes found: 12 of 12 (100%)" in report and "Rates matching the answer key: 9 of 9 (100%)" in report

    cli.main(["price", "--from", "Dallas, TX 75207", "--to", "Atlanta, GA", "--equipment", "dry van", "--lanes",
              str(case / "lanes.csv"), "--broker", str(case / "broker.toml"), "--date", "2026-09-21",
              "--customer", "Wide World Importers"])
    priced = capsys.readouterr().out
    assert "Median carrier pay $1,875 -> suggested linehaul $2,225 (16% margin), from same cities" in priced
    with pytest.raises(SystemExit, match="Could not read the place"):
        cli.main(["price", "--from", "Dallas", "--to", "Atlanta, GA", "--lanes", str(case / "lanes.csv")])
