"""Writes the prospect list and the draft: CSV and Excel for prospects, Word and Markdown for the letter."""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from agentkit.tables import Sheet, write_csv, write_xlsx

from .draft import Draft
from .funders import Funder
from .org import Org
from .prospect import Prospect

PROSPECT_COLUMNS = ["rank", "funder", "status", "fit", "best_program", "suggested_ask", "deadline", "application",
                    "why", "priorities", "concerns", "requirements", "ruled_out", "checks", "website", "contact"]


def _quotes(items) -> str:
    return " | ".join(f"{q.point}: \"{q.quote}\"" for q in items)


def prospect_rows(prospects: list[Prospect]) -> list[dict]:
    rows = []
    for n, p in enumerate(prospects, 1):
        f = p.funder
        rows.append({"rank": n if p.status != "ruled out" else "", "funder": f.name, "status": p.status,
                     "fit": "" if p.fit is None else p.fit, "best_program": p.best_program,
                     "suggested_ask": "" if p.suggested_ask is None else f"{p.suggested_ask:.0f}",
                     "deadline": "rolling" if f.rolling else (f.deadline.isoformat() if f.deadline else f.deadline_text),
                     "application": p.application or f.application, "why": p.why, "priorities": _quotes(p.priorities),
                     "concerns": _quotes(p.concerns), "requirements": _quotes(p.requirements),
                     "ruled_out": "; ".join(f.out), "checks": "; ".join([p.error] * bool(p.error) + p.checks + f.notes),
                     "website": f.website, "contact": f.contact})
    return rows


def write_prospects(out: Path, prospects: list[Prospect]) -> list[dict]:
    rows = prospect_rows(prospects)
    write_csv(out / "prospects.csv", rows, PROSPECT_COLUMNS)
    header = [c.replace("_", " ").capitalize() for c in PROSPECT_COLUMNS]
    yellow = {i for i, r in enumerate(rows) if r["checks"] and r["status"] != "ruled out"}
    write_xlsx(out / "prospects.xlsx", [Sheet("Prospects", header, [[r[c] for c in PROSPECT_COLUMNS] for r in rows],
                                              yellow, {2: 30, 9: 50, 10: 60, 11: 50, 12: 50, 13: 40, 14: 50})])
    (out / "prospects.json").write_text(json.dumps({"prospects": [
        {**r, "out": p.funder.out, "notes": p.funder.notes, "error": p.error} for r, p in zip(rows, prospects)]},
        indent=1))
    return rows


# --- the draft ----------------------------------------------------------------------------------------------------

def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:40]


def letter_parts(d: Draft, org: Org, funder: Funder) -> tuple[str, list[tuple[str, str]], str]:
    req = d.request
    salutation = req.salutation or (f"Dear {funder.contact}," if funder.contact else f"Dear {funder.name} team,")
    closing = "Sincerely,\n\n" + (req.signer or org.name)
    return salutation, [(spec.name, text) for spec, text, _ in d.sections], closing


def write_draft(out: Path, d: Draft, org: Org, funder: Funder, ev, today: date) -> dict:
    """draft_<funder>.docx and .md, review_draft.csv, and draft.json (for score)."""
    import docx
    from docx.shared import Pt

    name = f"draft_{slug(funder.name)}"
    req = d.request
    salutation, sections, closing = letter_parts(d, org, funder)
    title = f"{req.kind}: {req.program}"

    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"To: {funder.name}\nFrom: {org.name}\nDate: {today.strftime('%B %d, %Y').replace(' 0', ' ')}\n"
                      f"Request: ${req.amount:,.0f}")
    doc.add_paragraph(salutation)
    for heading, text in sections:
        doc.add_heading(heading, level=2)
        for para in [x for x in text.split("\n") if x.strip()]:
            doc.add_paragraph(para.strip())
    doc.add_paragraph(closing)
    doc.add_page_break()
    doc.add_heading("Checks and sources (delete this page before sending)", level=1)
    for f in d.flags:
        doc.add_paragraph(("MUST FIX: " if f.hold else "Check: ") + (f"{f.section}: " if f.section else "") + f.text,
                          style="List Bullet")
    for spec, _, ids in d.sections:
        doc.add_heading(spec.name, level=2)
        if not ids:
            doc.add_paragraph("No sources cited.")
        for i in ids:
            p = ev.items[i]
            doc.add_paragraph(f"[{i}] {p.source}: {p.text[:400]}", style="List Bullet")
    doc.save(str(out / f"{name}.docx"))

    md = [f"# {title}", "", f"To: {funder.name}  ", f"From: {org.name}  ", f"Request: ${req.amount:,.0f}", "",
          salutation, ""]
    for heading, text in sections:
        md += [f"## {heading}", "", text, ""]
    md += [closing.replace("\n", "  \n"), ""]
    (out / f"{name}.md").write_text("\n".join(md), encoding="utf-8")

    rows = [{"must_fix": "yes" if f.hold else "", "section": f.section, "check": f.text} for f in d.flags]
    rows.sort(key=lambda r: r["must_fix"] != "yes")
    write_csv(out / "review_draft.csv", rows, ["must_fix", "section", "check"])
    (out / "draft.json").write_text(json.dumps({
        "funder": funder.name, "program": req.program, "amount": req.amount, "ready": d.ready, "error": d.error,
        "words": d.words, "sections": [{"name": s.name, "text": t, "sources": ids, "words": len(t.split())}
                                       for s, t, ids in d.sections],
        "flags": [{"section": f.section, "hold": f.hold, "text": f.text} for f in d.flags]}, indent=1))
    return {"docx": out / f"{name}.docx", "md": out / f"{name}.md", "review_items": len(rows)}
