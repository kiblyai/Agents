"""Reads CSV and Excel tables into rows, and writes the CSV and Excel outputs the agents hand to a person."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

YELLOW = "FFF2CC"  # rows a person should check


def norm_header(h) -> str:
    """'Variant SKU', 'variant_sku' and ' VARIANT-SKU ' all become 'variant sku'."""
    return re.sub(r"[\s_\-]+", " ", str(h or "")).strip().lower()


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows_from_matrix(matrix: list[list], strip: bool = True) -> list[dict[str, str]]:
    """The first row with at least two filled cells is the header; empty rows are skipped."""
    start = next((i for i, r in enumerate(matrix) if sum(1 for v in r if _cell(v)) >= 2), None)
    if start is None:
        return []
    header = [norm_header(h) for h in matrix[start]]
    rows = []
    for r in matrix[start + 1:]:
        vals = [_cell(v) if strip or not isinstance(v, str) else v for v in r]
        if not any(_cell(v) for v in vals):
            continue
        rows.append({h: (vals[i] if i < len(vals) else "") for i, h in enumerate(header) if h})
    return rows


def read_table(path: str | Path, sheet: str | None = None, strip: bool = True) -> list[dict[str, str]]:
    """Rows of a CSV or Excel file as {normalized header: text}. Excel reads the named sheet, else the first.
    strip=False keeps spaces around values (for cleanup tools that report them)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".csv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return _rows_from_matrix(list(csv.reader(f)), strip)
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = _rows_from_matrix([list(r) for r in ws.iter_rows(values_only=True)], strip)
        wb.close()
        return rows
    if ext == ".xls":
        raise ValueError(f"{path.name}: old .xls files are not supported; open it in Excel and save as .xlsx")
    raise ValueError(f"{path.name}: expected a .csv or .xlsx file")


def pick(row: dict[str, str], *aliases: str) -> str:
    """The first non-empty value among the column names given (compared after norm_header)."""
    for a in aliases:
        v = row.get(norm_header(a), "")
        if v:
            return v
    return ""


def has_column(rows: list[dict[str, str]], *aliases: str) -> bool:
    keys = set(rows[0]) if rows else set()
    return any(norm_header(a) in keys for a in aliases)


def sheet_text(path_or_bytes) -> str:
    """An Excel workbook as plain text, one line per row ("A | B | C"), for a model or a text check to read."""
    import io

    from openpyxl import load_workbook

    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else str(path_or_bytes)
    wb = load_workbook(src, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        lines = []
        for r in ws.iter_rows(values_only=True):
            vals = [_cell(v) for v in r]
            while vals and not vals[-1]:
                vals.pop()
            if any(vals):  # empty cells inside a row stay, so the columns line up
                lines.append(" | ".join(vals))
        if lines and len(wb.worksheets) > 1:
            out.append(f"[Sheet: {ws.title}]")
        out += lines
    wb.close()
    return "\n".join(out)


def write_csv(path: str | Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


@dataclass
class Sheet:
    title: str
    header: list[str]
    rows: list[list] = field(default_factory=list)
    yellow: set[int] = field(default_factory=set)  # row numbers (0-based, after the header) to highlight
    widths: dict[int, float] = field(default_factory=dict)  # column number (1-based) -> width
    formats: dict[int, str] = field(default_factory=dict)  # column number (1-based) -> number format, e.g. '"$"#,##0'


def write_xlsx(path, sheets: list[Sheet]) -> None:
    """A workbook with a bold, frozen header row on each sheet and yellow rows to check."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor=YELLOW)
    for sh in sheets:
        ws = wb.create_sheet(sh.title[:31])
        ws.append(sh.header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, row in enumerate(sh.rows):
            ws.append(row)
            if i in sh.yellow:
                for c in ws[ws.max_row]:
                    c.fill = fill
        for col, w in sh.widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
                fmt = sh.formats.get(c.column)
                if fmt and isinstance(c.value, (int, float)):
                    c.number_format = fmt
        ws.freeze_panes = "A2"
    wb.save(path if hasattr(path, "write") else str(path))


def xlsx_bytes(sheets: list[Sheet]) -> bytes:
    """An Excel file whose bytes are the same on every run (fixed timestamps), for synthetic examples and tests."""
    import io
    import re as _re
    import zipfile

    raw = io.BytesIO()
    write_xlsx(raw, sheets)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-01-01T00:00:00Z", data)
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, data)
    return out.getvalue()
