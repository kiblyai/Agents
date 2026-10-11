"""Reads documents into labelled sections: PDF pages, Word and Markdown headings, plain text."""

from __future__ import annotations

import re
from pathlib import Path

TEXT_TYPES = {".md", ".txt", ".markdown"}
DOC_TYPES = TEXT_TYPES | {".pdf", ".docx"}


def chunk(text: str, max_words: int = 180) -> list[str]:
    """Split text into passages of up to max_words, keeping paragraphs together where possible."""
    paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n|\n(?=[-*•] )", text)]
    out, cur, n = [], [], 0
    for p in filter(None, paras):
        words = p.split()
        while len(words) > max_words:  # one very long paragraph
            if cur:
                out.append(" ".join(cur))
                cur, n = [], 0
            out.append(" ".join(words[:max_words]))
            words = words[max_words:]
        if n + len(words) > max_words and cur:
            out.append(" ".join(cur))
            cur, n = [], 0
        cur.append(" ".join(words))
        n += len(words)
    if cur:
        out.append(" ".join(cur))
    return out


def sections_markdown(text: str) -> list[tuple[str, str]]:
    """(heading, body) for each Markdown section; text before the first heading has heading ""."""
    sections, heading, buf = [], "", []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)", line)
        if m:
            if buf:
                sections.append((heading, "\n".join(buf)))
            heading, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if buf:
        sections.append((heading, "\n".join(buf)))
    return sections


def read_text_file(path: Path) -> list[tuple[str, str]]:
    return [(f"{path.name}" + (f" § {h}" if h else ""), body)
            for h, body in sections_markdown(path.read_text(encoding="utf-8", errors="replace"))]


def read_pdf(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [(f"{path.name} p.{i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]


def read_docx(path: Path) -> list[tuple[str, str]]:
    import docx

    d = docx.Document(str(path))
    sections, heading, buf = [], "", []
    for para in d.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading") and para.text.strip():
            if buf:
                sections.append((f"{path.name}" + (f" § {heading}" if heading else ""), "\n\n".join(buf)))
            heading, buf = para.text.strip(), []
        elif para.text.strip():
            buf.append(para.text.strip())
    if buf:
        sections.append((f"{path.name}" + (f" § {heading}" if heading else ""), "\n\n".join(buf)))
    for n, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells if c.text.strip()) for row in table.rows]
        text = "\n\n".join(r for r in rows if r)
        if text:
            sections.append((f"{path.name} table {n}", text))
    return sections


def read_document(path: str | Path) -> list[tuple[str, str]]:
    """(source label, text) sections of a PDF, Word, Markdown or text file. Raises ValueError for other types."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in TEXT_TYPES:
        return read_text_file(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    raise ValueError(f"{path.name}: unsupported type (use PDF, Word .docx, Markdown or text)")


def document_files(folder: str | Path, types: set[str] = DOC_TYPES) -> tuple[list[Path], list[str]]:
    """Readable documents in a folder (and its subfolders), in name order, plus the files skipped."""
    files, skipped = [], []
    for p in sorted(Path(folder).rglob("*")):
        if not p.is_file() or p.name.startswith((".", "~$")):
            continue
        if p.suffix.lower() in types:
            files.append(p)
        else:
            skipped.append(f"{p.name} (unsupported type)")
    return files, skipped
