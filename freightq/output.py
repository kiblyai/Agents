"""Writes the run: draft replies, the review list, a TMS import file for tenders, an Excel workbook and JSON."""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict
from pathlib import Path

from .broker import BLOCKING, EQUIPMENT_LABELS, FIELD_LABELS, BrokerConfig
from .extract import EmailResult, Shipment
from .pipeline import RunResult
from .reply import draft_reply, money

YELLOW = "FFF2CC"
REVIEW_COLUMNS = ["order", "email", "lane", "customer", "kind", "check", "reply"]
LOAD_COLUMNS = ["customer", "load_number", "references", "pickup_city", "pickup_state", "pickup_zip", "pickup_date",
                "pickup_time", "stops", "delivery_city", "delivery_state", "delivery_zip", "delivery_date",
                "delivery_time", "equipment", "weight_lbs", "commodity", "pieces", "temperature", "hazmat",
                "requirements", "customer_rate", "est_carrier_pay", "email", "needs_review", "notes"]


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_replies(run: RunResult, cfg: BrokerConfig, folder: Path) -> int:
    folder.mkdir(parents=True, exist_ok=True)
    for old in folder.glob("*.txt"):
        old.unlink()
    n = 0
    for r in run.results:
        text = draft_reply(r, cfg)
        if text:
            path = folder / (Path(r.email.file).stem + ".txt")
            path.write_text(text, encoding="utf-8")
            r.reply_file = f"{folder.name}/{path.name}"
            n += 1
    return n


def _notes(s: Shipment) -> str:
    return "; ".join(f.text for f in s.flags)


def review_rows(run: RunResult) -> list[dict]:
    """What a person should look at, most urgent first: tenders, then lanes to price or check, then the rest."""
    rows = []

    def add(order: int, r: EmailResult, lane: str, check: str) -> None:
        rows.append({"order": order, "email": r.email.file, "lane": lane, "customer": r.customer or r.email.sender,
                     "kind": r.kind.replace("_", " "), "check": check, "reply": r.reply_file})

    for r in run.results:
        if r.error:
            add(4, r, "", r.error)
            continue
        for f in r.flags:
            add(2, r, "", f.text)
        if r.kind == "other":
            add(4, r, "", "not a quote request or tender, so no reply was drafted; check nothing was missed")
            continue
        for s in r.shipments:
            lane = s.lane_text()
            q = s.quote
            if s.kind == "tender":
                decide = "tender: accept or decline"
                if s.rate is not None and q is not None and q.buy is not None:
                    decide += (f"; rate {money(s.rate)} vs recent carrier pay {money(q.buy)} "
                               f"({100 * q.margin:.0f}% margin)")
                add(1, r, lane, "; ".join([decide] + [f.text for f in s.flags]))
                continue
            blocked = [f for f in s.missing if f in BLOCKING]
            if s.flags:
                add(2, r, lane, _notes(s))
            elif blocked:
                add(3, r, lane, f"no rate: the reply asks for the {', '.join(FIELD_LABELS[f] for f in blocked)}")
            elif q is None or not q.priced:
                add(2, r, lane, "no rate: price by hand")
    rows.sort(key=lambda x: x["order"])
    return rows


def load_rows(run: RunResult) -> list[dict]:
    """Tenders as rows for a TMS import."""
    rows = []
    for r in run.results:
        for s in r.shipments:
            if s.kind != "tender":
                continue
            load_no = next((ref for ref in s.references if not re.match(r"\s*(?:PO|BOL|P/O)\b", ref, re.I)), "")
            q = s.quote
            rows.append({
                "customer": s.customer, "load_number": load_no, "references": "; ".join(s.references),
                "pickup_city": s.origin["city"], "pickup_state": s.origin["state"], "pickup_zip": s.origin["zip"],
                "pickup_date": s.pickup_date, "pickup_time": s.pickup_time,
                "stops": "; ".join(s.place_text(p) for p in s.stops),
                "delivery_city": s.destination["city"], "delivery_state": s.destination["state"],
                "delivery_zip": s.destination["zip"], "delivery_date": s.delivery_date,
                "delivery_time": s.delivery_time, "equipment": EQUIPMENT_LABELS.get(s.equipment, ""),
                "weight_lbs": "" if s.weight_lbs is None else f"{s.weight_lbs:.0f}", "commodity": s.commodity,
                "pieces": s.pieces, "temperature": s.temperature, "hazmat": "yes" if s.hazmat else "no",
                "requirements": "; ".join(s.requirements),
                "customer_rate": "" if s.rate is None else f"{s.rate:.2f}",
                "est_carrier_pay": "" if q is None or q.buy is None else f"{q.buy:.2f}",
                "email": s.email, "needs_review": "yes", "notes": _notes(s),
            })
    return rows


def _as_dict(r: EmailResult) -> dict:
    e = r.email
    return {"file": e.file, "sent": e.sent.isoformat() if e.sent else "", "sender": e.sender, "subject": e.subject,
            "kind": r.kind, "customer": r.customer, "contact": r.contact, "reply_file": r.reply_file,
            "error": r.error, "flags": [asdict(f) for f in r.flags],
            "shipments": [asdict(s) | {"priced": bool(s.quote and s.quote.priced)} for s in r.shipments]}


def write_json(path: Path, run: RunResult) -> None:
    path.write_text(json.dumps({"emails": [_as_dict(r) for r in run.results]}, indent=1, default=str))


def _sheet(wb, title: str, header: list[str], rows: list[list], yellow: set[int], widths: dict[int, float]):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    fill = PatternFill("solid", fgColor=YELLOW)
    for i, row in enumerate(rows):
        ws.append(row)
        if i in yellow:
            for c in ws[ws.max_row]:
                c.fill = fill
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 40)
    ws.freeze_panes = "A2"
    return ws


def write_xlsx(path: Path, run: RunResult) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    ships = [s for r in run.results for s in r.shipments]
    header = ["Email", "Kind", "Customer", "Lane", "Origin", "Stops", "Destination", "Pickup", "Delivery",
              "Equipment", "Weight (lbs)", "Commodity", "Pieces", "Temp", "Hazmat", "Requirements", "References",
              "Trucks", "Customer's rate", "Est. carrier pay", "Linehaul", "Extras", "All-in rate", "Margin",
              "Priced from", "Missing", "Check before sending"]
    rows, yellow = [], set()
    for i, s in enumerate(ships):
        q = s.quote
        priced = q is not None and q.priced
        rows.append([
            s.email, s.kind.replace("_", " "), s.customer, s.lane, s.place_text(s.origin),
            "; ".join(s.place_text(p) for p in s.stops), s.place_text(s.destination),
            " ".join(x for x in (s.pickup_date, s.pickup_time) if x),
            " ".join(x for x in (s.delivery_date, s.delivery_time) if x),
            EQUIPMENT_LABELS.get(s.equipment, ""), s.weight_lbs, s.commodity, s.pieces, s.temperature,
            "yes" if s.hazmat else "", "; ".join(s.requirements), "; ".join(s.references), s.truckloads, s.rate,
            q.buy if q else None, q.linehaul if priced else None,
            "; ".join(f"{k} {v:,.0f}" for k, v in q.extras) if priced else "",
            q.total if priced else None, round(q.margin, 3) if priced and q.margin is not None else None,
            (f"{q.basis}, {len(q.comparables)} loads" if q and q.basis else "no history") if q else "",
            ", ".join(s.missing), _notes(s)])
        if s.flags or not priced or s.kind == "tender":
            yellow.add(i)
    ws = _sheet(wb, "Shipments", header, rows, yellow, {1: 24, 3: 20, 5: 22, 6: 18, 7: 22, 12: 18, 27: 60})
    for row in ws.iter_rows(min_row=2):
        for c in (row[18], row[19], row[20], row[22]):
            c.number_format = '"$"#,##0'
        row[23].number_format = "0%"

    comp = []
    for s in ships:
        for x in (s.quote.comparables if s.quote else []):
            comp.append([s.email, s.lane, x.day, x.customer, x.lane_text(), EQUIPMENT_LABELS.get(x.equipment, ""),
                         x.miles, x.carrier_pay, x.customer_rate, x.row])
    _sheet(wb, "Comparable loads", ["Email", "Lane", "Date", "Customer", "Lane history", "Equipment", "Miles",
                                    "Carrier pay", "Customer rate", "CSV row"], comp, set(), {1: 24, 4: 20, 5: 44})

    em = [[r.email.file, r.email.sent.strftime("%Y-%m-%d %H:%M") if r.email.sent else "", r.email.sender,
           r.email.subject, r.kind.replace("_", " "), len(r.shipments), r.reply_file,
           "; ".join([r.error] + [f.text for f in r.flags]).strip("; ")] for r in run.results]
    _sheet(wb, "Emails", ["File", "Sent", "From", "Subject", "Kind", "Lanes", "Reply", "Notes"], em,
           {i for i, r in enumerate(run.results) if r.error or r.flags}, {1: 24, 3: 30, 4: 40, 7: 24, 8: 50})
    review = review_rows(run)
    _sheet(wb, "Review", [c.capitalize() for c in REVIEW_COLUMNS], [[x[c] for c in REVIEW_COLUMNS] for x in review],
           set(), {2: 24, 3: 40, 4: 20, 6: 80, 7: 24})
    wb.save(path)


def summary(run: RunResult, usage, lanes_used: int, lane_problems: list[str]) -> dict:
    ships = [s for r in run.results for s in r.shipments]
    kinds = {k: sum(1 for r in run.results if r.kind == k and not r.error) for k in ("quote_request", "tender", "other")}
    return {
        "emails": len(run.results), "quote_requests": kinds["quote_request"], "tenders": kinds["tender"],
        "other": kinds["other"], "not_processed": sum(1 for r in run.results if r.error),
        "lanes": len(ships), "lanes_priced": sum(1 for s in ships if s.quote and s.quote.priced),
        "lanes_flagged": sum(1 for s in ships if s.flags), "lanes_missing_info": sum(1 for s in ships if s.missing),
        "replies_drafted": sum(1 for r in run.results if r.reply_file),
        "review_items": len(review_rows(run)), "lane_history_loads": lanes_used,
        "lane_history_rows_skipped": lane_problems, "stopped_reason": run.stopped_reason,
        "llm_requests": usage.requests, "llm_retries": usage.retries, "llm_repairs": usage.repairs,
        "llm_failed_attempts": usage.failed_attempts, "prompt_tokens": usage.prompt_tokens,
        "completion_tokens": usage.completion_tokens, "cost_usd": round(usage.cost_usd, 6),
        "models_used": usage.models, "from_cache": run.from_cache,
    }
