"""Writes the chronology: Word (for the law firm), Excel and CSV (for checking), JSON (for scoring), PDF bookmarks."""

from __future__ import annotations

import csv
import json
from dataclasses import asdict
from datetime import date
from pathlib import Path

from .chronology import REVIEW_COLUMNS, Chronology, page_range
from .extract import Entry
from .records import CaseFile

YELLOW = "FFF2CC"


def us_date(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%m/%d/%Y") if iso else "undated"


def money(x: float | None) -> str:
    return "" if x is None else f"${x:,.2f}"


def pages_text(e: Entry) -> str:
    return page_range(e.pages) + (f" (copy: {page_range(e.duplicate_pages)})" if e.duplicate_pages else "")


def sources(pages: list[int], case: CaseFile) -> str:
    """Pages as "file p.1-2; other.pdf p.4"."""
    by_file: dict[str, list[int]] = {}
    for n in pages:
        if 1 <= n <= len(case.pages):
            p = case.page(n)
            by_file.setdefault(p.file, []).append(p.file_page)
    return "; ".join(f"{f} p.{page_range(ps)}" for f, ps in by_file.items())


def bookmarks(c: Chronology) -> list[tuple[str, int]]:
    out = []
    for e in c.entries:
        if e.pages:
            who = e.provider or e.facility
            out.append((f"{us_date(e.date)} {e.visit_type}" + (f" - {who}" if who else "") +
                        (" [check]" if e.needs_review else ""), e.pages[0]))
    return out


def write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_json(path: Path, c: Chronology) -> None:
    data = {"entries": [asdict(e) for e in c.entries], "bills": [asdict(b) for b in c.bills],
            "gaps": [asdict(g) for g in c.gaps], "providers": c.providers, "bill_totals": c.bill_totals,
            "total_billed": c.total_billed, "days_to_first_treatment": c.days_to_first_treatment}
    path.write_text(json.dumps(data, indent=1))


# --- Excel --------------------------------------------------------------------------------------------------------

def write_xlsx(path: Path, c: Chronology, case: CaseFile) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    fill = PatternFill("solid", start_color=YELLOW, end_color=YELLOW)

    def sheet(title: str, header: list[str], rows: list[list], widths: list[int], flag_col: int | None = None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
            if flag_col is not None and row[flag_col]:
                for cell in ws[ws.max_row]:
                    cell.fill = fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, date):
                    cell.number_format = "mm/dd/yyyy"
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    sheet("Chronology",
          ["Date", "Before injury", "Provider", "Facility", "Visit type", "Complaints", "Findings", "Diagnoses",
           "Treatment", "Work status", "Pages", "Duplicate copies", "Source", "Confidence", "Needs review",
           "Review note"],
          [[date.fromisoformat(e.date) if e.date else "undated", "yes" if e.pre_incident else "", e.provider,
            e.facility, e.visit_type, e.complaints, e.findings, e.diagnosis_text(), e.treatment, e.work_status,
            page_range(e.pages), page_range(e.duplicate_pages), sources(e.pages, case), e.confidence,
            "yes" if e.needs_review else "", e.review_note] for e in c.entries],
          [11, 8, 20, 22, 18, 30, 30, 30, 30, 20, 9, 9, 22, 10, 8, 40], flag_col=14)
    sheet("Treatment gaps", ["From", "To", "Days", "Last visit before", "Next visit after"],
          [[date.fromisoformat(g.start), date.fromisoformat(g.end), g.days, g.before, g.after] for g in c.gaps],
          [11, 11, 7, 40, 40])
    sheet("Bills", ["Date of service", "Provider", "Code", "Description", "Charge", "Pages", "Needs review",
                    "Review note"],
          [[date.fromisoformat(b.date) if b.date else "", b.provider, b.code, b.description, b.charge,
            page_range(b.pages), "yes" if b.needs_review else "", b.review_note] for b in c.bills],
          [11, 28, 8, 36, 11, 8, 8, 40], flag_col=6)
    totals = sheet("Bill totals", ["Provider", "Charges", "Total", "Charges to check"],
                   [[t["provider"], t["lines"], t["total"], t["to_check"]] for t in c.bill_totals] +
                   [["Total", sum(t["lines"] for t in c.bill_totals), c.total_billed,
                     sum(t["to_check"] for t in c.bill_totals)]], [36, 9, 13, 10])
    totals.cell(totals.max_row, 1).font = Font(bold=True)
    for ws in (wb["Bills"], totals):
        col = 5 if ws.title == "Bills" else 3
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = '"$"#,##0.00'
    sheet("Review", [k.capitalize() for k in REVIEW_COLUMNS], [[r[k] for k in REVIEW_COLUMNS] for r in c.review],
          [11, 11, 10, 40, 60])
    sheet("Pages", ["Page", "File", "File page", "Characters", "Used by", "Check"],
          [[p["page"], p["file"], p["file_page"], p["characters"], p["used_by"], p["check"]] for p in c.pages],
          [7, 30, 8, 10, 30, 50])
    wb.save(path)


# --- Word -----------------------------------------------------------------------------------------------------------

def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float], shade_rows: set[int] = frozenset()):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.autofit = False
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
        _shade(cell, "D9E2F3")
    repeat = OxmlElement("w:tblHeader")  # repeat the header row on every printed page
    repeat.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            if isinstance(value, list):  # [(label, text), ...] as separate lines with bold labels
                cell.text = ""
                first = True
                for label, text in value:
                    para = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    if label:
                        para.add_run(f"{label}: ").bold = True
                    para.add_run(text)
            else:
                cell.text = str(value)
            if i in shade_rows:
                _shade(cell, YELLOW)
    for col, w in zip(t.columns, widths):  # Word reads the cell widths, LibreOffice the column grid
        col.width = Inches(w)
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)
    return t


def write_docx(path: Path, c: Chronology, case: CaseFile, *, patient: str, doi: date | None,
               combined_name: str = "records_combined.pdf") -> None:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    doc.styles["Normal"].font.size = Pt(9)

    doc.add_heading("Medical chronology" + (f": {patient}" if patient else ""), level=0)
    files = len({p.file for p in case.pages})
    doc.add_paragraph(" · ".join(filter(None, [
        f"Date of injury: {doi.strftime('%m/%d/%Y')}" if doi else "",
        f"Records reviewed: {len(case.pages)} pages from {files} file{'s' if files != 1 else ''}",
        f"Prepared {date.today().strftime('%m/%d/%Y')}"])))
    to_check = sum(1 for r in c.review if r["kind"] in ("entry", "bill"))
    if to_check:
        p = doc.add_paragraph()
        p.add_run(f"DRAFT: {to_check} yellow row{'s' if to_check != 1 else ''} to check. Fix them, then delete the "
                  "\"Check before sending\" column and this line.").italic = True

    doc.add_heading("Summary", level=1)
    dated = [e for e in c.entries if e.date and not e.pre_incident]
    bullets = []
    if dated:
        bullets.append(f"Treatment from {us_date(dated[0].date)} to {us_date(dated[-1].date)}: {len(dated)} "
                       f"encounters with {len(c.providers)} providers.")
    if c.days_to_first_treatment is not None:
        bullets.append(f"First treatment {c.days_to_first_treatment} days after the injury.")
    if c.gaps:
        bullets.append("Treatment gaps over the limit: " + "; ".join(
            f"{us_date(g.start)} to {us_date(g.end)} ({g.days} days)" for g in c.gaps) + ".")
    else:
        bullets.append("No treatment gaps over the limit.")
    pre = [e for e in c.entries if e.pre_incident]
    if pre:
        bullets.append(f"{len(pre)} record{'s' if len(pre) != 1 else ''} from before the injury (marked "
                       f"\"Before injury\"): " + "; ".join(f"{us_date(e.date)} {e.visit_type}" for e in pre) + ".")
    if c.bills:
        bullets.append(f"Total billed: {money(c.total_billed)} across {len(c.bills)} charges.")
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    if c.providers:
        _table(doc, ["Provider", "First visit", "Last visit", "Visits"],
               [[r["provider"], us_date(r["first"]) if r["first"] else "", us_date(r["last"]) if r["last"] else "",
                 str(r["visits"])] for r in c.providers], [5.0, 1.5, 1.5, 1.0])

    doc.add_heading("Chronology", level=1)
    rows, shade = [], set()
    for i, e in enumerate(c.entries):
        details = [(label, text) for label, text in (
            ("Complaints", e.complaints), ("Findings", e.findings), ("Diagnoses", e.diagnosis_text()),
            ("Treatment", e.treatment), ("Work status", e.work_status)) if text]
        rows.append([us_date(e.date) + ("\nBefore injury" if e.pre_incident else ""),
                     "\n".join(x for x in (e.provider, e.facility) if x), e.visit_type, details or [("", "")],
                     pages_text(e), e.review_note if e.needs_review else ""])
        if e.needs_review:
            shade.add(i)
    _table(doc, ["Date", "Provider", "Visit", "Details", "Pages", "Check before sending"], rows,
           [0.95, 1.55, 1.2, 4.3, 0.8, 1.0], shade)

    doc.add_heading("Treatment gaps", level=1)
    if c.gaps:
        _table(doc, ["From", "To", "Days", "Last visit before", "Next visit after"],
               [[us_date(g.start), us_date(g.end), str(g.days), g.before, g.after] for g in c.gaps],
               [1.0, 1.0, 0.6, 3.6, 3.6])
    else:
        doc.add_paragraph("None over the limit.")

    doc.add_heading("Bills", level=1)
    if c.bill_totals:
        _table(doc, ["Provider", "Charges", "Total billed"],
               [[t["provider"], str(t["lines"]), money(t["total"])] for t in c.bill_totals] +
               [["Total", str(len(c.bills)), money(c.total_billed)]], [6.0, 1.3, 1.5])
        doc.add_paragraph("Totals are sums of the itemized charges, which are listed in the Excel file (Bills sheet).")
    else:
        doc.add_paragraph("No bills in the records.")

    doc.add_paragraph()
    doc.add_paragraph(f"Page numbers refer to the stamped page numbers in {combined_name}, which also has a bookmark "
                      "for every entry.").italic = True
    doc.save(path)


def summary(run, c: Chronology, case: CaseFile, usage) -> dict:
    return {
        "pages": len(case.pages),
        "pages_without_text": sum(1 for p in case.pages if p.needs_ocr),
        "files_read": case.files_read,
        "files_skipped": case.files_skipped,
        "entries": len(c.entries),
        "entries_need_review": sum(1 for e in c.entries if e.needs_review),
        "entries_before_injury": sum(1 for e in c.entries if e.pre_incident),
        "duplicate_copies_merged": sum(1 for e in c.entries if e.duplicate_pages),
        "treatment_gaps": len(c.gaps),
        "days_to_first_treatment": c.days_to_first_treatment,
        "bill_lines": len(c.bills),
        "bill_lines_need_review": sum(1 for b in c.bills if b.needs_review),
        "total_billed": c.total_billed,
        "review_items": len(c.review),
        "pages_not_processed": len(run.not_processed),
        "stopped_early": run.stopped_reason,
        "batch_errors": run.batch_errors,
        "batches": run.batches,
        "batches_from_cache": run.batches_from_cache,
        "llm_requests": usage.requests,
        "llm_failed_attempts": usage.failed_attempts,
        "llm_repairs": usage.repairs,
        "prompt_tokens": usage.prompt_tokens,
        "completion_tokens": usage.completion_tokens,
        "cost_usd": round(usage.cost_usd, 4),
        "models_used": usage.models,
    }
