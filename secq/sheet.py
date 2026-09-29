"""Finds the questions in a questionnaire spreadsheet and writes draft answers back into a copy of it."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

ID_RE = re.compile(r"^(?:question |control |item )?(?:id|#|no\.?|number|ref(?:erence)?|s/?n)$")
QUESTION_WORDS = ("question", "requirement", "query", "criteria", "control")
ANSWER_WORDS = ("answer", "response", "yes/no", "yes / no", "y/n", "compliant", "compliance", "implemented")
COMMENT_WORDS = ("comment", "explanation", "notes", "details", "justification", "evidence", "additional information",
                 "description of")
# Words that start a question even without a question mark ("Describe your backup process").
QUESTION_STARTERS = {"describe", "provide", "list", "explain", "detail", "specify", "confirm", "indicate", "state",
                     "do", "does", "did", "is", "are", "was", "were", "has", "have", "can", "will", "would", "should",
                     "how", "what", "when", "where", "which", "who", "why"}
REVIEW_FILL = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")


@dataclass
class Layout:
    sheet: str
    header_row: int
    q_col: int
    a_col: int | None = None
    c_col: int | None = None
    id_col: int | None = None

    def describe(self) -> str:
        def col(c):
            return get_column_letter(c) if c else "-"
        return (f"sheet '{self.sheet}': header row {self.header_row}, question col {col(self.q_col)}, "
                f"answer col {col(self.a_col)}, comment col {col(self.c_col)}, id col {col(self.id_col)}")


@dataclass
class Question:
    key: str  # "Sheet!row"
    sheet: str
    row: int
    qid: str
    text: str
    section: str = ""
    prefilled: bool = False


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def classify_header(h: str) -> str | None:
    if not h:
        return None
    if ID_RE.match(h):
        return "id"
    if any(w in h for w in COMMENT_WORDS):
        return "comment"
    if any(w in h for w in ANSWER_WORDS):
        return "answer"
    if any(w in h for w in QUESTION_WORDS):
        return "question"
    return None


def detect(ws, max_scan: int = 25, max_cols: int = 40) -> Layout | None:
    for r in range(1, min(ws.max_row, max_scan) + 1):
        found: dict[str, list[tuple[int, str]]] = {}
        for c in range(1, min(ws.max_column, max_cols) + 1):
            h = _norm(ws.cell(r, c).value)
            kind = classify_header(h)
            if kind:
                found.setdefault(kind, []).append((c, h))
        if "question" in found and ("answer" in found or "comment" in found):
            q = next((c for c, h in found["question"] if "question" in h), found["question"][0][0])
            first = lambda k: found[k][0][0] if k in found else None  # noqa: E731
            return Layout(sheet=ws.title, header_row=r, q_col=q, a_col=first("answer"), c_col=first("comment"),
                          id_col=first("id"))
    return None


def open_workbook(path: str | Path) -> Workbook:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        wb = Workbook()
        ws = wb.active
        ws.title = path.stem[:31] or "Sheet1"
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.reader(f):
                ws.append(row)
        return wb
    if path.suffix.lower() == ".xls":
        raise ValueError("Old .xls files are not supported; open it in Excel and save as .xlsx")
    return load_workbook(str(path))


def _is_section(text: str, qid: str, bold: bool, answered: bool) -> bool:
    if qid or answered or "?" in text:
        return False
    first = text.split()[0].lower().strip(".:") if text.split() else ""
    if first in QUESTION_STARTERS:
        return False
    return bold or len(text.split()) <= 4


def read_questions(wb: Workbook, sheet: str | None = None, override: dict | None = None
                   ) -> tuple[dict[str, Layout], list[Question]]:
    """Return ({sheet: layout}, questions) for every sheet whose layout can be found (or the one requested)."""
    layouts: dict[str, Layout] = {}
    questions: list[Question] = []
    names = [sheet] if sheet else wb.sheetnames
    for name in names:
        if name not in wb.sheetnames:
            raise ValueError(f"no sheet named '{name}'; sheets are: {', '.join(wb.sheetnames)}")
        ws = wb[name]
        lay = detect(ws)
        if override:
            base = lay or Layout(sheet=name, header_row=1, q_col=1)
            for k, v in override.items():
                if v is not None:
                    setattr(base, k, v)
            lay = base
        if lay is None:
            continue
        layouts[name] = lay
        section = ""
        for r in range(lay.header_row + 1, ws.max_row + 1):
            cell = ws.cell(r, lay.q_col)
            text = re.sub(r"\s+", " ", str(cell.value or "")).strip()
            if len(text) < 3:
                continue
            qid = str(ws.cell(r, lay.id_col).value or "").strip() if lay.id_col else ""
            answered = any(ws.cell(r, c).value not in (None, "") for c in (lay.a_col, lay.c_col) if c)
            bold = bool(cell.font and cell.font.b)
            if _is_section(text, qid, bold, answered):
                section = text
                continue
            questions.append(Question(key=f"{name}!{r}", sheet=name, row=r, qid=qid, text=text, section=section,
                                      prefilled=answered))
    return layouts, questions


def column_number(letter: str | None) -> int | None:
    return column_index_from_string(letter.strip().upper()) if letter else None


EXTRA_HEADERS = ["Draft sources", "Draft confidence", "Review note"]


def write_draft(src: str | Path, out_path: str | Path, layouts: dict[str, Layout], results: list) -> None:
    """Write answers into a copy of the questionnaire. Rows needing review are highlighted yellow.

    Three helper columns are added at the right (delete them before sending the questionnaire back).
    """
    wb = open_workbook(src)
    extra: dict[str, int] = {}
    answer_col: dict[str, int] = {}
    for name, lay in layouts.items():
        ws = wb[name]
        col = ws.max_column + 1
        if not lay.a_col and not lay.c_col:  # no answer column at all: add one
            ws.cell(lay.header_row, col, "Draft answer")
            answer_col[name] = col
            col += 1
        for i, h in enumerate(EXTRA_HEADERS):
            ws.cell(lay.header_row, col + i, h)
        extra[name] = col
    for r in results:
        if r.sheet not in layouts:
            continue
        lay, ws, start = layouts[r.sheet], wb[r.sheet], extra[r.sheet]
        if lay.a_col and lay.c_col:
            targets = [(lay.a_col, r.answer), (lay.c_col, r.explanation)]
        else:
            col = lay.a_col or lay.c_col or answer_col[r.sheet]
            joined = ". ".join(p.rstrip(".") for p in (r.answer, r.explanation) if p)
            targets = [(col, (joined + ".") if joined else "")]
        for col, value in targets:
            cell = ws.cell(r.row, col)
            if value:
                cell.value = value
            if r.needs_review:
                cell.fill = REVIEW_FILL
        ws.cell(r.row, start, "; ".join(r.source_labels))
        ws.cell(r.row, start + 1, r.confidence)
        ws.cell(r.row, start + 2, r.review_note)
    wb.save(str(out_path))
