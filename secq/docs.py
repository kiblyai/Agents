"""Reads a company's security documents into short citable passages, plus previously approved answers."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

TEXT_TYPES = {".md", ".txt", ".markdown"}
SHEET_TYPES = {".xlsx", ".xlsm", ".csv"}


@dataclass
class Passage:
    id: str
    source: str  # human-readable location, e.g. "security-policy.pdf p.3" or "access.md § MFA"
    text: str


@dataclass
class LibraryAnswer:
    """A question the company has answered before (from a past questionnaire)."""

    id: str
    question: str
    answer: str
    source: str


@dataclass
class KnowledgeBase:
    passages: list[Passage] = field(default_factory=list)
    library: list[LibraryAnswer] = field(default_factory=list)
    files_read: list[str] = field(default_factory=list)
    files_skipped: list[str] = field(default_factory=list)


def chunk(text: str, max_words: int = 180) -> list[str]:
    """Split text into passages of up to max_words, keeping paragraphs together where possible."""
    paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n|\n(?=[-*•] )", text)]
    out, cur, n = [], [], 0
    for p in filter(None, paras):
        words = p.split()
        while len(words) > max_words:  # one very long paragraph
            if cur:
                out.append(" ".join(cur))
                cur, n = [], 0
            out.append(" ".join(words[:max_words]))
            words = words[max_words:]
        if n + len(words) > max_words and cur:
            out.append(" ".join(cur))
            cur, n = [], 0
        cur.append(" ".join(words))
        n += len(words)
    if cur:
        out.append(" ".join(cur))
    return out


def _sections_markdown(text: str) -> list[tuple[str, str]]:
    sections, heading, buf = [], "", []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)", line)
        if m:
            if buf:
                sections.append((heading, "\n".join(buf)))
            heading, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if buf:
        sections.append((heading, "\n".join(buf)))
    return sections


def read_text_file(path: Path) -> list[tuple[str, str]]:
    return [(f"{path.name}" + (f" § {h}" if h else ""), body)
            for h, body in _sections_markdown(path.read_text(encoding="utf-8", errors="replace"))]


def read_pdf(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [(f"{path.name} p.{i}", page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]


def read_docx(path: Path) -> list[tuple[str, str]]:
    import docx

    d = docx.Document(str(path))
    sections, heading, buf = [], "", []
    for para in d.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        if style.lower().startswith("heading") and para.text.strip():
            if buf:
                sections.append((f"{path.name}" + (f" § {heading}" if heading else ""), "\n\n".join(buf)))
            heading, buf = para.text.strip(), []
        elif para.text.strip():
            buf.append(para.text.strip())
    if buf:
        sections.append((f"{path.name}" + (f" § {heading}" if heading else ""), "\n\n".join(buf)))
    for n, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells if c.text.strip()) for row in table.rows]
        text = "\n\n".join(r for r in rows if r)
        if text:
            sections.append((f"{path.name} table {n}", text))
    return sections


def read_library(path: Path, start_id: int) -> list[LibraryAnswer]:
    """Question/answer pairs from a past questionnaire (xlsx or csv)."""
    from .sheet import open_workbook, read_questions

    wb = open_workbook(path)
    layouts, questions = read_questions(wb)
    out = []
    for q in questions:
        lay = layouts[q.sheet]
        ws = wb[q.sheet]
        parts = [ws.cell(q.row, c).value for c in (lay.a_col, lay.c_col) if c]
        vals = [str(p).strip() for p in parts if str(p or "").strip()]
        # "Yes" + "Encrypted with AES-256." reads as "Yes. Encrypted with AES-256."
        answer = ". ".join(v.rstrip(".") for v in vals) + "." if len(vals) > 1 else "".join(vals)
        if answer:
            out.append(LibraryAnswer(id=f"L{start_id + len(out)}", question=q.text, answer=answer,
                                     source=f"{path.name} {q.sheet} row {q.row}"))
    return out


def load_kb(folder: str | Path, max_words: int = 180) -> KnowledgeBase:
    kb = KnowledgeBase()
    files = sorted(p for p in Path(folder).rglob("*") if p.is_file() and not p.name.startswith((".", "~$")))
    for path in files:
        ext = path.suffix.lower()
        try:
            if ext in SHEET_TYPES:
                pairs = read_library(path, len(kb.library) + 1)
                if not pairs:
                    kb.files_skipped.append(f"{path.name} (no answered questions found)")
                    continue
                kb.library.extend(pairs)
                kb.files_read.append(f"{path.name} ({len(pairs)} past answers)")
                continue
            if ext in TEXT_TYPES:
                sections = read_text_file(path)
            elif ext == ".pdf":
                sections = read_pdf(path)
            elif ext == ".docx":
                sections = read_docx(path)
            else:
                kb.files_skipped.append(f"{path.name} (unsupported type)")
                continue
        except Exception as e:  # a corrupt file should not stop the run
            kb.files_skipped.append(f"{path.name} ({type(e).__name__}: {e})")
            continue
        before = len(kb.passages)
        for source, body in sections:
            for text in chunk(body, max_words):
                kb.passages.append(Passage(id=f"P{len(kb.passages) + 1}", source=source, text=text))
        if len(kb.passages) == before:
            kb.files_skipped.append(f"{path.name} (no readable text; scanned PDF?)")
        else:
            kb.files_read.append(f"{path.name} ({len(kb.passages) - before} passages)")
    return kb


def write_csv_rows(path: Path, rows: list[dict], columns: list[str]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
